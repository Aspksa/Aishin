from __future__ import annotations

import base64
import json
import sys
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from docx import Document as SmokeDocxDocument
from openpyxl import Workbook as SmokeWorkbook
from pypdf import PdfWriter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def _configure_utf8_output() -> None:
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if callable(reconfigure):
            reconfigure(encoding="utf-8", errors="replace")


_configure_utf8_output()

from fastapi.testclient import TestClient

from app.db import connect
from app.core.memory import MemoryCandidate
from app.main import app, engine


def main() -> int:
    checks: dict[str, object] = {}

    try:
        with TestClient(app) as client:
            home = client.get("/")
            if home.status_code != 200:
                raise RuntimeError(
                    f"/ returned HTTP {home.status_code}: {home.text[:300]}"
                )
            required_ui = (
                'id="chat-form"',
                'id="messages"',
                'id="simple-mode"',
                'id="simple-learning"',
                'class="technical-brain"',
                'id="brain-flow"',
                'id="pending-decisions"',
                'id="settings-module"',
                'id="cloud-api-key"',
                'id="brain-links"',
                'data-brain-node="reflection"',
                'data-brain-node="learning-plan"',
                'data-brain-node="experiment"',
                'data-brain-node="learning-check"',
                'data-brain-node="consolidation"',
                'id="reflection-quality"',
                'id="learning-plan-stream"',
                'id="experiment-stream"',
                'id="context-budget-value"',
                'id="development-open"',
                'id="development-score"',
                'id="development-module"',
                'id="development-components"',
                'id="development-chart"',
                'id="development-reasons"',
                'id="dev-count-knowledge"',
                '/static/intelligence.css',
                '/static/intelligence.js',
                '/static/proactive_intelligence.css',
                '/static/proactive_intelligence.js',
                'id="attention-module"',
                'id="proactive-awareness-score"',
                'id="proactive-incidents"',
                'id="proactive-calibration"',
                '/static/evolution.css',
                '/static/evolution.js',
                'id="evolution-module"',
                'id="evolution-score"',
                'id="evolution-capabilities"',
                'id="evolution-variants"',
                'id="evolution-curriculum"',
                '/static/research.css',
                '/static/research.js',
                'id="research-module"',
                'id="research-score"',
                'id="research-gaps"',
                'id="research-evidence"',
                'id="research-claims"',
                'id="research-contradictions"',
                '/static/communication.css',
                '/static/communication.js',
                'id="communication-module"',
                'id="communication-score"',
                'id="communication-turns"',
                'id="communication-skills"',
                'id="communication-runtime"',
                'id="communication-prefs"',
                'id="communication-events"',
                '/static/documents.css',
                '/static/documents.js',
                'id="documents-module"',
                'id="documents-score"',
                'id="documents-list"',
                'id="documents-facts"',
                'id="documents-conflicts"',
                'id="documents-upload-btn"',
                'id="scope-select"',
                'value="project:aishin"',
                '/static/app.js?v=0.0.15',
                '/static/live_brain.js?v=0.0.15',
                'role="dialog"',
                'aria-modal="true"',
            )
            missing_ui = [
                marker for marker in required_ui
                if marker not in home.text
            ]
            if missing_ui:
                raise RuntimeError(
                    f"Новая UI-компоновка неполная: {missing_ui}"
                )
            checks["ui_layout"] = {
                "status": "ok",
                "chat_first": True,
                "technical_brain_collapsed": True,
            }

            cloud_settings = client.get("/api/settings/cloudru")
            if cloud_settings.status_code != 200:
                raise RuntimeError(
                    f"/api/settings/cloudru returned HTTP "
                    f"{cloud_settings.status_code}"
                )
            cloud_data = cloud_settings.json()
            if "api_key" in cloud_data or "AISHIN_CLOUDRU_API_KEY" in cloud_data:
                raise RuntimeError("Cloud settings API leaked API key field")

            update_status = client.get("/api/system/update/status")
            if update_status.status_code != 200:
                raise RuntimeError(
                    f"/api/system/update/status returned HTTP "
                    f"{update_status.status_code}"
                )
            if update_status.json().get("mode") not in {"git", "zip"}:
                raise RuntimeError("Updater returned unknown mode")

            checks["cloud_settings_ui"] = {
                "status": "ok",
                "configured": cloud_data.get("configured"),
                "secret_not_exposed": True,
                "update_mode": update_status.json().get("mode"),
            }

            health = client.get("/health")
            if health.status_code != 200:
                raise RuntimeError(
                    f"/health returned HTTP {health.status_code}: {health.text[:300]}"
                )
            health_data = health.json()
            if health_data.get("version") != "0.0.15":
                raise RuntimeError(
                    "Health должен сообщать Aishin Core 0.0.15"
                )
            with connect() as conn:
                schema_row = conn.execute(
                    "SELECT MAX(version) AS version FROM schema_migrations"
                ).fetchone()
            schema_version = int(schema_row["version"] or 0)
            if schema_version != 26:
                raise RuntimeError(
                    f"Ожидалась database schema 26, получено {schema_version}"
                )
            checks["health"] = {
                **health_data,
                "schema_version": schema_version,
            }

            state = client.get("/api/assistant/state")
            if state.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/state returned HTTP {state.status_code}: "
                    f"{state.text[:300]}"
                )
            state_data = state.json()
            checks["state"] = {
                "identity": state_data.get("identity", {}).get("name"),
                "runtime_status": state_data.get("state", {}).get("status"),
                "profile_source": state_data.get("identity", {}).get(
                    "profile_source"
                ),
                "schema_present": bool(state_data.get("state")),
            }

            scoped_state = client.get(
                "/api/assistant/state",
                params={"scope": "project:aishin"},
            )
            if scoped_state.status_code != 200:
                raise RuntimeError("Scoped state endpoint недоступен")
            if (
                scoped_state.json().get("state", {}).get("current_scope")
                != "project:aishin"
            ):
                raise RuntimeError(
                    "UI scope должен доходить до engine.snapshot без смешивания"
                )
            checks["scope_isolation"] = {
                "status": "ok",
                "project_scope": "project:aishin",
            }

            engine.events.emit(
                "smoke.scope.personal",
                scope="personal",
                payload={"sentinel": "personal"},
                importance=0.1,
            )
            engine.events.emit(
                "smoke.scope.project",
                scope="project:aishin",
                payload={"sentinel": "project"},
                importance=0.1,
            )
            personal_memory_change = engine.memory.consolidate(
                MemoryCandidate(
                    content=(
                        "Runtime smoke: personal scope isolation sentinel."
                    ),
                    scope="personal",
                    kind="smoke_scope",
                    confidence=1.0,
                    importance=0.1,
                ),
                source="runtime_smoke",
            )
            project_memory_change = engine.memory.consolidate(
                MemoryCandidate(
                    content=(
                        "Runtime smoke: project Aishin scope isolation "
                        "sentinel."
                    ),
                    scope="project:aishin",
                    kind="smoke_scope",
                    confidence=1.0,
                    importance=0.1,
                ),
                source="runtime_smoke",
            )
            if (
                personal_memory_change.memory_id is None
                or project_memory_change.memory_id is None
            ):
                raise RuntimeError(
                    "Не удалось создать scoped memory sentinels"
                )

            project_state_data = scoped_state.json()
            project_recent_events = (
                project_state_data.get("recent_events") or []
            )
            if any(
                item.get("scope") != "project:aishin"
                for item in project_recent_events
            ):
                raise RuntimeError(
                    "engine.snapshot смешивает recent_events между scope"
                )

            project_changes = client.get(
                "/api/assistant/memory-changes",
                params={
                    "scope": "project:aishin",
                    "limit": 30,
                },
            )
            personal_changes = client.get(
                "/api/assistant/memory-changes",
                params={
                    "scope": "personal",
                    "limit": 30,
                },
            )
            if (
                project_changes.status_code != 200
                or personal_changes.status_code != 200
            ):
                raise RuntimeError(
                    "Scoped memory-changes endpoint недоступен"
                )
            if not project_changes.json() or not personal_changes.json():
                raise RuntimeError(
                    "Scoped memory-changes не вернул sentinels"
                )
            if any(
                item.get("scope") != "project:aishin"
                for item in project_changes.json()
            ):
                raise RuntimeError(
                    "Project memory_changes смешаны с другим scope"
                )
            if any(
                item.get("scope") != "personal"
                for item in personal_changes.json()
            ):
                raise RuntimeError(
                    "Personal memory_changes смешаны с другим scope"
                )

            scoped_brain = client.get(
                "/api/assistant/brain",
                params={"scope": "project:aishin"},
            )
            if (
                scoped_brain.status_code != 200
                or scoped_brain.json().get("state", {}).get(
                    "current_scope"
                ) != "project:aishin"
            ):
                raise RuntimeError(
                    "Legacy brain endpoint должен использовать explicit scope"
                )

            baseline_interactions = int(
                engine.state.load().interaction_count
            )
            atomic_workers = 12
            with ThreadPoolExecutor(
                max_workers=atomic_workers
            ) as pool:
                futures = [
                    pool.submit(
                        engine.state.interaction,
                        "runtime_smoke",
                        scope="personal",
                    )
                    for _ in range(atomic_workers)
                ]
                for future in futures:
                    future.result()
            state_after_threads = engine.state.load()
            if state_after_threads.interaction_count != (
                baseline_interactions + atomic_workers
            ):
                raise RuntimeError(
                    "StateManager потерял interaction_count при "
                    "параллельных обновлениях"
                )
            if state_after_threads.current_scope != "personal":
                raise RuntimeError(
                    "Atomic interaction не сохранил переданный scope"
                )
            if tuple(engine.continuous_learning.scopes) != tuple(
                engine.RUNTIME_SCOPES
            ):
                raise RuntimeError(
                    "Continuous Learning не изолирован по runtime scopes"
                )

            checks["scope_isolation"].update(
                {
                    "events_scoped": True,
                    "memory_changes_scoped": True,
                    "state_atomic": True,
                    "background_scopes": list(
                        engine.RUNTIME_SCOPES
                    ),
                }
            )

            tools = client.get("/api/assistant/tools")
            if tools.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/tools returned HTTP {tools.status_code}"
                )
            tools_catalog = client.get("/api/assistant/tools")
            if tools_catalog.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/tools returned HTTP {tools_catalog.status_code}"
                )
            tool_names = {
                item.get("name")
                for item in tools_catalog.json()
            }
            if "project.rollback_write" not in tool_names:
                raise RuntimeError("Rollback tool отсутствует в Tool Registry")

            rollback_missing = client.post(
                "/api/assistant/execution/attempts/999999999/rollback",
                json={"scope": "personal", "approved": False},
            )
            if rollback_missing.status_code not in {400, 403}:
                raise RuntimeError(
                    "Rollback API должен отклонять неизвестный или нелокальный вызов"
                )
            checks["atomic_write_rollback"] = {
                "status": "ok",
                "rollback_tool_present": True,
                "rollback_api_present": True,
            }

            checks["tools"] = {
                "status": "ok",
                "count": len(tools.json()),
            }

            bridge_scope = "smoke:execution-bridge"
            read_selection = SimpleNamespace(
                selected=SimpleNamespace(
                    matched_tool="project.read_text",
                )
            )
            read_bridge = engine.action_execution.prepare(
                scope=bridge_scope,
                query="прочитать файл README.md",
                selection=read_selection,
                decision_quality=0.95,
                unresolved_count=0,
            )
            if (
                read_bridge.state != "executed_read_only"
                or not read_bridge.executed
                or read_bridge.result.get("status") != "success"
            ):
                raise RuntimeError(
                    "Action Selection -> ToolRegistry read-only bridge "
                    "не выполнился безопасно"
                )

            write_selection = SimpleNamespace(
                selected=SimpleNamespace(
                    matched_tool="project.write_text",
                )
            )
            write_bridge = engine.action_execution.prepare(
                scope=bridge_scope,
                query="изменить файл README.md",
                selection=write_selection,
                decision_quality=0.95,
                unresolved_count=0,
            )
            if write_bridge.executed or write_bridge.state != (
                "arguments_unresolved"
            ):
                raise RuntimeError(
                    "Execution bridge не должен додумывать mutating arguments"
                )

            flow_scope = "smoke:brain-flow"
            flow_request = "smoke-flow-1"
            engine.brain_flow.begin(
                request_id=flow_request,
                scope=flow_scope,
                intent="verification",
            )
            engine.brain_flow.phase(
                request_id=flow_request,
                scope=flow_scope,
                phase="documents",
                detail={"chunks": 2},
            )
            flow_running = engine.brain_flow.snapshot(scope=flow_scope)
            executing = [
                node
                for node in flow_running.get("nodes") or []
                if node.get("status") == "executing"
            ]
            if (
                flow_running.get("status") != "running"
                or len(executing) != 1
                or executing[0].get("id") != "documents"
            ):
                raise RuntimeError(
                    "BrainFlow должен показывать только реально выполняемую фазу"
                )
            engine.brain_flow.phase(
                request_id=flow_request,
                scope=flow_scope,
                phase="metacognition",
            )
            engine.brain_flow.finish(
                request_id=flow_request,
                scope=flow_scope,
            )
            flow_done = engine.brain_flow.snapshot(scope=flow_scope)
            if flow_done.get("status") != "completed":
                raise RuntimeError(
                    "BrainFlow не зафиксировал завершение запроса"
                )
            if len(flow_done.get("nodes") or []) != 24:
                raise RuntimeError("BrainFlow node contract нарушен")

            checks["brain_wiring"] = {
                "status": "ok",
                "read_only_execution": True,
                "mutating_guess_blocked": True,
                "real_time_nodes": 24,
            }

            meta_scope = "smoke:metacognition-independence"
            same_document_meta = engine.metacognition.assess(
                scope=meta_scope,
                intent="verification",
                recalled_memories=[],
                semantic_used=False,
                planner_notices=[],
                sensor_readings=[],
                graph_stats={},
                external_evidence=[
                    {
                        "source": "document",
                        "source_type": "document",
                        "source_group": "document:101",
                        "confidence": 0.98,
                        "independence": 1.0,
                    },
                    {
                        "source": "document",
                        "source_type": "document",
                        "source_group": "document:101",
                        "confidence": 0.97,
                        "independence": 1.0,
                    },
                ],
            )
            if same_document_meta.status == "confident":
                raise RuntimeError(
                    "Два chunks одного документа не должны давать "
                    "high-stakes confident"
                )
            if not any(
                "недостаточно независимых первичных" in item
                for item in same_document_meta.missing_data
            ):
                raise RuntimeError(
                    "Metacognition должна видеть один document source group"
                )

            two_document_meta = engine.metacognition.assess(
                scope=meta_scope,
                intent="verification",
                recalled_memories=[],
                semantic_used=False,
                planner_notices=[],
                sensor_readings=[],
                graph_stats={},
                external_evidence=[
                    {
                        "source": "document",
                        "source_type": "document",
                        "source_group": "document:201",
                        "confidence": 0.98,
                        "independence": 1.0,
                    },
                    {
                        "source": "document",
                        "source_type": "document",
                        "source_group": "document:202",
                        "confidence": 0.97,
                        "independence": 1.0,
                    },
                ],
            )
            if any(
                "недостаточно независимых первичных" in item
                for item in two_document_meta.missing_data
            ):
                raise RuntimeError(
                    "Два независимых документа должны удовлетворять "
                    "source-group boundary"
                )

            duplicated_memory_meta = engine.metacognition.assess(
                scope=meta_scope,
                intent="action",
                recalled_memories=[
                    {
                        "id": 1,
                        "confidence": 1.0,
                        "retrieval_score": 1.0,
                        "kind": "fact",
                        "tags": [],
                    }
                ],
                semantic_used=True,
                planner_notices=[],
                sensor_readings=[],
                graph_stats={},
                external_evidence=[
                    {
                        "source": "research",
                        "source_type": "semantic_memory",
                        "source_ref": "1",
                        "source_group": "memory",
                        "confidence": 1.0,
                        "independence": 1.0,
                    }
                ],
            )
            if duplicated_memory_meta.status == "confident":
                raise RuntimeError(
                    "Одна память через lexical+semantic не должна "
                    "считаться двумя независимыми high-stakes опорами"
                )

            grounding_scope = "smoke:response-grounding"
            supported_grounding = engine.response_grounding.assess(
                request_id="grounding-supported",
                scope=grounding_scope,
                mode="VERIFY",
                response="Договор действует до 31 декабря 2026 года.",
                evidence=[
                    {
                        "source": "document",
                        "source_type": "document",
                        "source_group": "document:501",
                        "document_id": 501,
                        "chunk_id": 1,
                        "confidence": 1.0,
                        "content": "Договор действует до 31 декабря 2026 года.",
                        "provenance": {
                            "document_id": 501,
                            "chunk_id": 1,
                            "page": 2,
                        },
                    }
                ],
            )
            if (
                not supported_grounding.applicable
                or supported_grounding.status != "strong"
                or supported_grounding.claims_supported != 1
                or supported_grounding.claims_unsupported != 0
            ):
                raise RuntimeError(
                    "Response Grounding не распознал прямую документальную опору"
                )

            numeric_mismatch = engine.response_grounding.assess(
                request_id="grounding-number-mismatch",
                scope=grounding_scope,
                mode="VERIFY",
                response="Стоимость договора составляет 999999 рублей.",
                evidence=[
                    {
                        "source": "document",
                        "source_type": "document",
                        "source_group": "document:502",
                        "document_id": 502,
                        "chunk_id": 1,
                        "confidence": 1.0,
                        "content": "Стоимость договора составляет 100 рублей.",
                        "provenance": {"document_id": 502, "chunk_id": 1},
                    }
                ],
            )
            if (
                numeric_mismatch.claims_supported != 0
                or numeric_mismatch.claims_unsupported < 1
            ):
                raise RuntimeError(
                    "Response Grounding не должен подтверждать несовпадающее число"
                )

            negation_mismatch = engine.response_grounding.assess(
                request_id="grounding-negation-mismatch",
                scope=grounding_scope,
                mode="VERIFY",
                response="Договор не действует после 2026 года.",
                evidence=[
                    {
                        "source": "document",
                        "source_type": "document",
                        "source_group": "document:503",
                        "document_id": 503,
                        "chunk_id": 1,
                        "confidence": 1.0,
                        "content": "Договор действует после 2026 года.",
                        "provenance": {"document_id": 503, "chunk_id": 1},
                    }
                ],
            )
            if negation_mismatch.claims_supported != 0:
                raise RuntimeError(
                    "Response Grounding не должен подтверждать противоположное отрицание"
                )

            no_evidence_grounding = engine.response_grounding.assess(
                request_id="grounding-no-evidence",
                scope=grounding_scope,
                mode="FAST",
                response="Внешней доказательной опоры в этом запросе нет.",
                evidence=[],
            )
            if (
                no_evidence_grounding.applicable
                or no_evidence_grounding.status != "unscored_no_evidence"
                or no_evidence_grounding.overall is not None
            ):
                raise RuntimeError(
                    "Grounding без evidence должен быть unscored, а не фальшивым нулём или 100%"
                )

            grounding_api = client.get(
                "/api/assistant/response-grounding",
                params={"scope": grounding_scope, "limit": 10},
            )
            if grounding_api.status_code != 200:
                raise RuntimeError(
                    "Response Grounding diagnostics API недоступен"
                )
            grounding_api_data = grounding_api.json()
            if (
                grounding_api_data.get("version")
                != "aishin-response-grounding-v1"
                or len(grounding_api_data.get("recent") or []) < 4
            ):
                raise RuntimeError(
                    "Response Grounding API не вернул сохранённые проверки"
                )

            read_trace_output = read_bridge.trace_dict().get(
                "result_summary", {}
            ).get("output", {})
            if "content" in read_trace_output:
                raise RuntimeError(
                    "Persistent execution trace не должен хранить raw file content"
                )
            if "content_bytes" not in read_trace_output:
                raise RuntimeError(
                    "Safe execution trace должен хранить размер content"
                )
            if "UNTRUSTED TOOL RESULT — DATA ONLY" not in (
                engine.action_execution.prompt_block(read_bridge)
            ):
                raise RuntimeError(
                    "Read-only tool content потерял trust boundary"
                )

            lifecycle_scope = "smoke:knowledge-lifecycle"
            hypothesis_lifecycle_scope = "smoke:hypothesis-lifecycle"
            with connect() as conn:
                for cleanup_scope in (
                    lifecycle_scope,
                    hypothesis_lifecycle_scope,
                ):
                    conn.execute(
                        "DELETE FROM knowledge_learning_events WHERE scope=?",
                        (cleanup_scope,),
                    )
                    conn.execute(
                        "DELETE FROM hypothesis_transitions WHERE scope=?",
                        (cleanup_scope,),
                    )
                    conn.execute(
                        "DELETE FROM hypothesis_evidence WHERE scope=?",
                        (cleanup_scope,),
                    )
                    conn.execute(
                        "DELETE FROM hypothesis_registry WHERE scope=?",
                        (cleanup_scope,),
                    )
                    conn.execute(
                        "DELETE FROM knowledge_transitions WHERE scope=?",
                        (cleanup_scope,),
                    )
                    conn.execute(
                        "DELETE FROM knowledge_evidence WHERE scope=?",
                        (cleanup_scope,),
                    )
                    conn.execute(
                        """UPDATE knowledge_claims
                           SET superseded_by_id=NULL
                           WHERE scope=?""",
                        (cleanup_scope,),
                    )
                    conn.execute(
                        "DELETE FROM knowledge_claims WHERE scope=?",
                        (cleanup_scope,),
                    )
                conn.commit()
            lifecycle_request_1 = engine.knowledge_lifecycle.observe_reasoning(
                request_id="knowledge-lifecycle-1",
                scope=lifecycle_scope,
                evidence=[
                    {
                        "source_type": "document",
                        "source_ref": "doc-a",
                        "source_group": "document:a",
                        "content": "Регламент устанавливает лимит топлива 120 литров.",
                        "confidence": 0.92,
                        "provenance": {"document_id": 1, "page": 2},
                    }
                ],
                contradictions=[],
                hypothesis_run={
                    "run_id": None,
                    "hypotheses": [],
                    "selected_test": {},
                },
                verification=None,
            )
            first_claims = engine.knowledge_lifecycle.claims(
                scope=lifecycle_scope,
                limit=10,
            )
            if (
                len(first_claims) != 1
                or first_claims[0].get("state") != "observed"
            ):
                raise RuntimeError(
                    "Один источник не должен автоматически становиться verified knowledge"
                )

            lifecycle_request_2 = engine.knowledge_lifecycle.observe_reasoning(
                request_id="knowledge-lifecycle-2",
                scope=lifecycle_scope,
                evidence=[
                    {
                        "source_type": "document",
                        "source_ref": "doc-a",
                        "source_group": "document:a",
                        "content": "Регламент устанавливает лимит топлива 120 литров.",
                        "confidence": 0.92,
                    },
                    {
                        "source_type": "research",
                        "source_ref": "manual-b",
                        "source_group": "manual:b",
                        "content": "Регламент устанавливает лимит топлива 120 литров.",
                        "confidence": 0.88,
                    },
                ],
                contradictions=[],
                hypothesis_run={
                    "run_id": None,
                    "hypotheses": [],
                    "selected_test": {},
                },
                verification={
                    "ran": True,
                    "unresolved": [],
                    "consistency": {"conflicts": []},
                },
            )
            verified_claims = engine.knowledge_lifecycle.claims(
                scope=lifecycle_scope,
                state="verified",
                limit=10,
            )
            if len(verified_claims) != 1:
                raise RuntimeError(
                    "Два независимых evidence + clean Verification должны "
                    "поднять claim в verified"
                )

            engine.knowledge_lifecycle.observe_reasoning(
                request_id="knowledge-lifecycle-3",
                scope=lifecycle_scope,
                evidence=[
                    {
                        "source_type": "document",
                        "source_ref": "doc-a",
                        "source_group": "document:a",
                        "content": "Регламент устанавливает лимит топлива 120 литров.",
                        "confidence": 0.92,
                    }
                ],
                contradictions=[
                    {
                        "id": 99,
                        "source": "document",
                        "source_group": "document:c",
                        "summary": "Регламент не устанавливает лимит топлива 120 литров.",
                        "severity": 0.95,
                    }
                ],
                hypothesis_run={
                    "run_id": None,
                    "hypotheses": [],
                    "selected_test": {},
                },
                verification=None,
            )
            contradicted = engine.knowledge_lifecycle.claims(
                scope=lifecycle_scope,
                state="contradicted",
                limit=10,
            )
            if len(contradicted) != 1:
                raise RuntimeError(
                    "Сильное связанное противоречие должно переводить claim "
                    "в contradicted"
                )

            engine.knowledge_lifecycle.observe_reasoning(
                request_id="knowledge-lifecycle-4",
                scope=lifecycle_scope,
                evidence=[
                    {
                        "source_type": "document",
                        "source_ref": "doc-new-a",
                        "source_group": "document:new-a",
                        "content": "Обновлённый регламент устанавливает лимит топлива 130 литров.",
                        "confidence": 0.93,
                    },
                    {
                        "source_type": "research",
                        "source_ref": "manual-new-b",
                        "source_group": "manual:new-b",
                        "content": "Обновлённый регламент устанавливает лимит топлива 130 литров.",
                        "confidence": 0.89,
                    },
                ],
                contradictions=[],
                hypothesis_run={
                    "run_id": None,
                    "hypotheses": [],
                    "selected_test": {},
                },
                verification={
                    "ran": True,
                    "unresolved": [],
                    "consistency": {"conflicts": []},
                },
            )
            replacement_claims = [
                item
                for item in engine.knowledge_lifecycle.claims(
                    scope=lifecycle_scope,
                    state="verified",
                    limit=20,
                )
                if "130" in str(item.get("statement") or "")
            ]
            if len(replacement_claims) != 1:
                raise RuntimeError(
                    "Replacement knowledge must be verified before supersession"
                )
            superseded = engine.knowledge_lifecycle.supersede_claim(
                scope=lifecycle_scope,
                old_claim_id=int(contradicted[0]["id"]),
                new_claim_id=int(replacement_claims[0]["id"]),
                reason="new_regulation_replaces_previous_limit",
            )
            if (
                superseded.get("state") != "superseded"
                or int(superseded.get("superseded_by_id") or 0)
                != int(replacement_claims[0]["id"])
            ):
                raise RuntimeError(
                    "Knowledge supersession must preserve replacement linkage"
                )

            hypothesis_run = {
                "run_id": None,
                "hypotheses": [
                    {
                        "key": "fuel-limit-cause",
                        "title": "Лимит задан действующим регламентом",
                        "confidence": 0.74,
                        "supporting": [
                            "Лимит задан действующим регламентом 120 литров"
                        ],
                        "opposing": [],
                        "status": "leading",
                    }
                ],
                "selected_test": {
                    "competing_hypotheses": ["fuel-limit-cause"],
                },
            }
            engine.knowledge_lifecycle.observe_reasoning(
                request_id="hypothesis-lifecycle-1",
                scope=hypothesis_lifecycle_scope,
                evidence=[
                    {
                        "source_type": "document",
                        "source_ref": "d1",
                        "source_group": "document:d1",
                        "content": "Лимит задан действующим регламентом 120 литров",
                        "confidence": 0.90,
                    }
                ],
                contradictions=[],
                hypothesis_run=hypothesis_run,
                verification=None,
            )
            hypothesis_rows = engine.knowledge_lifecycle.hypotheses(
                scope=hypothesis_lifecycle_scope,
                limit=10,
            )
            if (
                len(hypothesis_rows) != 1
                or hypothesis_rows[0].get("state") == "confirmed"
            ):
                raise RuntimeError(
                    "Высокий confidence без независимой evidence не должен "
                    "подтверждать гипотезу"
                )

            confirmed_hypothesis_run = {
                "run_id": None,
                "hypotheses": [
                    {
                        "key": "fuel-limit-cause",
                        "title": "Лимит задан действующим регламентом",
                        "confidence": 0.78,
                        "supporting": [
                            "Источник A подтверждает действующий регламент 120 литров",
                            "Источник B подтверждает действующий регламент 120 литров",
                            "Источник C подтверждает действующий регламент 120 литров",
                        ],
                        "opposing": [],
                        "status": "leading",
                    }
                ],
                "selected_test": {},
            }
            engine.knowledge_lifecycle.observe_reasoning(
                request_id="hypothesis-lifecycle-2",
                scope=hypothesis_lifecycle_scope,
                evidence=[
                    {
                        "source_type": "document",
                        "source_ref": "hd1",
                        "source_group": "document:hd1",
                        "content": "Источник A подтверждает действующий регламент 120 литров",
                        "confidence": 0.92,
                    },
                    {
                        "source_type": "research",
                        "source_ref": "hr2",
                        "source_group": "research:hr2",
                        "content": "Источник B подтверждает действующий регламент 120 литров",
                        "confidence": 0.88,
                    },
                    {
                        "source_type": "memory",
                        "source_ref": "hm3",
                        "source_group": "memory:hm3",
                        "content": "Источник C подтверждает действующий регламент 120 литров",
                        "confidence": 0.86,
                    },
                ],
                contradictions=[],
                hypothesis_run=confirmed_hypothesis_run,
                verification={
                    "ran": True,
                    "unresolved": [],
                    "consistency": {"conflicts": []},
                },
            )
            confirmed_hypotheses = engine.knowledge_lifecycle.hypotheses(
                scope=hypothesis_lifecycle_scope,
                state="confirmed",
                limit=10,
            )
            if len(confirmed_hypotheses) != 1:
                raise RuntimeError(
                    "Three independent support groups + clean Verification "
                    "must allow hypothesis confirmation"
                )

            learning_cycle = engine.continuous_learning.run_cycle(
                scope=lifecycle_scope,
            )
            lifecycle_patterns = [
                item
                for item in engine.continuous_learning.patterns(
                    scope=lifecycle_scope,
                    limit=100,
                )
                if item.get("category") == "knowledge_lifecycle"
            ]
            if not any(
                item.get("pattern_key") == "knowledge_verified"
                for item in lifecycle_patterns
            ):
                raise RuntimeError(
                    "Continuous Learning did not ingest evidence-backed "
                    "knowledge lifecycle events"
                )

            knowledge_api = client.get(
                "/api/assistant/knowledge-lifecycle",
                params={"scope": lifecycle_scope, "limit": 20},
            )
            if knowledge_api.status_code != 200:
                raise RuntimeError("Knowledge Lifecycle API недоступен")
            lifecycle_api_data = knowledge_api.json()
            if (
                lifecycle_api_data.get("summary", {}).get("version")
                != "aishin-knowledge-lifecycle-v1"
            ):
                raise RuntimeError("Knowledge Lifecycle API version mismatch")
            if not lifecycle_api_data.get("transitions"):
                raise RuntimeError(
                    "Knowledge Lifecycle должен сохранять audit transitions"
                )
            if not any(
                isinstance(item.get("evidence"), list)
                and item.get("evidence")
                and item["evidence"][0].get("source_group")
                for item in lifecycle_api_data.get("claims") or []
            ):
                raise RuntimeError(
                    "Knowledge claims API должен возвращать evidence provenance"
                )
            hypothesis_provenance_api = client.get(
                "/api/assistant/knowledge-lifecycle/hypotheses",
                params={
                    "scope": hypothesis_lifecycle_scope,
                    "state": "confirmed",
                    "limit": 10,
                },
            )
            if hypothesis_provenance_api.status_code != 200:
                raise RuntimeError(
                    "Hypothesis Lifecycle provenance API недоступен"
                )
            hypothesis_api_rows = hypothesis_provenance_api.json()
            if not any(
                item.get("state") == "confirmed"
                and len(item.get("evidence") or []) >= 3
                for item in hypothesis_api_rows
            ):
                raise RuntimeError(
                    "Confirmed hypothesis должен раскрывать independent evidence"
                )
            supersede_api = client.post(
                (
                    "/api/assistant/knowledge-lifecycle/claims/"
                    f"{int(contradicted[0]['id'])}/supersede"
                ),
                json={
                    "scope": lifecycle_scope,
                    "new_claim_id": int(replacement_claims[0]["id"]),
                    "reason": "runtime_idempotent_supersession_check",
                },
            )
            if supersede_api.status_code != 403:
                raise RuntimeError(
                    "Knowledge supersession HTTP API должен сохранять "
                    "local-only guard в TestClient"
                )

            project_lifecycle = client.get(
                "/api/assistant/knowledge-lifecycle",
                params={"scope": "project:aishin", "limit": 10},
            )
            if project_lifecycle.status_code != 200:
                raise RuntimeError("Project-scoped Knowledge Lifecycle недоступен")
            if (
                project_lifecycle.json().get("summary", {}).get("scope")
                != "project:aishin"
            ):
                raise RuntimeError(
                    "Knowledge Lifecycle должен соблюдать scope isolation"
                )

            canonical_scope = "smoke:canonical-facts"
            with connect() as conn:
                conn.execute(
                    "DELETE FROM canonical_fact_events WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM canonical_fact_links WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM canonical_fact_evidence WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM canonical_fact_values WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM canonical_facts WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM knowledge_learning_events WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM knowledge_transitions WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM knowledge_evidence WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    """UPDATE knowledge_claims
                       SET superseded_by_id=NULL
                       WHERE scope=?""",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM knowledge_claims WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM relations WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM graph_changes WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM entities WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM memories WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM research_contradictions WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM research_claims WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM research_evidence WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM research_sessions WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM research_gaps WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM research_cycles WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM research_events WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM research_sources WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM research_state WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM document_contradictions WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM document_facts WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM document_ingestion_runs WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM document_events WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM document_chunk_vectors WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM document_chunks WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM document_sections WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM document_pages WHERE scope=?",
                    (canonical_scope,),
                )
                conn.execute(
                    "DELETE FROM documents WHERE scope=?",
                    (canonical_scope,),
                )

                def insert_document(
                    *,
                    sha: str,
                    filename: str,
                    family: str,
                    rank: int,
                    previous: int | None = None,
                ) -> int:
                    cur = conn.execute(
                        """INSERT INTO documents(
                               scope, sha256, filename, media_type, extension,
                               size_bytes, storage_path, status, parser,
                               parser_version, document_type, family_key,
                               version_label, version_rank, previous_version_id,
                               quality_score, extraction_coverage, studied_at
                           ) VALUES (
                               ?, ?, ?, 'text/plain', '.txt', 100, ?,
                               'studied', 'runtime_smoke', '1',
                               'regulation', ?, ?, ?, ?, 0.96, 1.0,
                               CURRENT_TIMESTAMP
                           )""",
                        (
                            canonical_scope,
                            sha,
                            filename,
                            f"smoke/{filename}",
                            family,
                            f"v{rank}",
                            rank,
                            previous,
                        ),
                    )
                    return int(cur.lastrowid)

                def insert_fact(
                    *,
                    document_id: int,
                    fact_key: str,
                    subject: str,
                    value: str,
                    normalized: str,
                    fact_type: str = "key_value",
                    confidence: float = 0.95,
                ) -> int:
                    cur = conn.execute(
                        """INSERT INTO document_facts(
                               scope, document_id, fact_key, subject,
                               predicate, value, normalized_value, fact_type,
                               confidence, status, provenance_json
                           ) VALUES (
                               ?, ?, ?, ?, 'has_value', ?, ?, ?, ?,
                               'grounded', ?
                           )""",
                        (
                            canonical_scope,
                            document_id,
                            fact_key,
                            subject,
                            value,
                            normalized,
                            fact_type,
                            confidence,
                            json.dumps(
                                {
                                    "runtime_smoke": True,
                                    "document_id": document_id,
                                },
                                ensure_ascii=False,
                            ),
                        ),
                    )
                    return int(cur.lastrowid)

                power_a = insert_document(
                    sha="a" * 64,
                    filename="power-a.txt",
                    family="power-source-a",
                    rank=1,
                )
                power_b = insert_document(
                    sha="b" * 64,
                    filename="power-b.txt",
                    family="power-source-b",
                    rank=1,
                )
                insert_fact(
                    document_id=power_a,
                    fact_key="engine_power:has_value",
                    subject="Engine power",
                    value="100 kW",
                    normalized="100 kw",
                    confidence=0.94,
                )
                insert_fact(
                    document_id=power_b,
                    fact_key="engine_power:has_value",
                    subject="Engine power",
                    value="100 kW",
                    normalized="100 kw",
                    confidence=0.92,
                )

                limit_old = insert_document(
                    sha="c" * 64,
                    filename="fuel-limit-v1.txt",
                    family="fuel-regulation",
                    rank=1,
                )
                limit_new = insert_document(
                    sha="d" * 64,
                    filename="fuel-limit-v2.txt",
                    family="fuel-regulation",
                    rank=2,
                    previous=limit_old,
                )
                insert_fact(
                    document_id=limit_old,
                    fact_key="fuel_limit:has_value",
                    subject="Fuel limit",
                    value="120 liters",
                    normalized="120 liters",
                    confidence=0.95,
                )
                insert_fact(
                    document_id=limit_new,
                    fact_key="fuel_limit:has_value",
                    subject="Fuel limit",
                    value="130 liters",
                    normalized="130 liters",
                    confidence=0.95,
                )

                interval_a = insert_document(
                    sha="e" * 64,
                    filename="interval-a.txt",
                    family="interval-source-a",
                    rank=1,
                )
                interval_b = insert_document(
                    sha="f" * 64,
                    filename="interval-b.txt",
                    family="interval-source-b",
                    rank=1,
                )
                insert_fact(
                    document_id=interval_a,
                    fact_key="service_interval:has_value",
                    subject="Service interval",
                    value="10000 km",
                    normalized="10000 km",
                    confidence=0.96,
                )
                insert_fact(
                    document_id=interval_b,
                    fact_key="service_interval:has_value",
                    subject="Service interval",
                    value="15000 km",
                    normalized="15000 km",
                    confidence=0.96,
                )

                meta_a = insert_document(
                    sha="1" * 64,
                    filename="meta-a.txt",
                    family="meta-a",
                    rank=1,
                )
                meta_b = insert_document(
                    sha="2" * 64,
                    filename="meta-b.txt",
                    family="meta-b",
                    rank=1,
                )
                insert_fact(
                    document_id=meta_a,
                    fact_key="document:mentions_date",
                    subject="document",
                    value="2026-01-01",
                    normalized="2026-01-01",
                    fact_type="date",
                    confidence=0.94,
                )
                insert_fact(
                    document_id=meta_b,
                    fact_key="document:mentions_date",
                    subject="document",
                    value="2026-02-01",
                    normalized="2026-02-01",
                    fact_type="date",
                    confidence=0.94,
                )

                research_session = int(
                    conn.execute(
                        """INSERT INTO research_sessions(
                               scope, question, trigger, status,
                               evidence_count, independent_groups,
                               claim_count, completed_at
                           ) VALUES (
                               ?, 'Oil specification', 'runtime_smoke',
                               'completed', 2, 2, 1, CURRENT_TIMESTAMP
                           )""",
                        (canonical_scope,),
                    ).lastrowid
                )
                research_evidence_ids = []
                for idx, group in enumerate(("manual:a", "manual:b"), 1):
                    evidence_id = int(
                        conn.execute(
                            """INSERT INTO research_evidence(
                                   scope, session_id, evidence_key,
                                   source_type, source_ref, source_group,
                                   title, content, stance, reliability,
                                   relevance, freshness, independence,
                                   evidence_score, content_hash, metadata_json
                               ) VALUES (
                                   ?, ?, ?, 'manual', ?, ?,
                                   'Oil manual', 'Oil specification is 5W-30',
                                   'support', 0.95, 0.95, 1.0, 1.0,
                                   0.93, ?, '{}'
                               )""",
                            (
                                canonical_scope,
                                research_session,
                                f"runtime-evidence-{idx}",
                                f"manual-{idx}",
                                group,
                                f"hash-{idx}",
                            ),
                        ).lastrowid
                    )
                    research_evidence_ids.append(evidence_id)

                research_claim_id = int(
                    conn.execute(
                        """INSERT INTO research_claims(
                               scope, claim_key, session_id, statement,
                               status, confidence, weighted_support,
                               support_count, independent_groups,
                               source_diversity, support_evidence_json
                           ) VALUES (
                               ?, 'oil_spec_claim', ?,
                               'Oil specification is 5W-30',
                               'trusted', 0.93, 0.93, 2, 2, 1, ?
                           )""",
                        (
                            canonical_scope,
                            research_session,
                            json.dumps(research_evidence_ids),
                        ),
                    ).lastrowid
                )
                promoted_memory_id = int(
                    conn.execute(
                        """INSERT INTO memories(
                               scope, kind, content, confidence, importance,
                               tags_json, source, fingerprint, memory_key,
                               status
                           ) VALUES (
                               ?, 'researched_knowledge',
                               'Oil specification is 5W-30',
                               0.93, 0.8, '["research"]', ?, ?,
                               'research_claim:oil_spec_claim', 'active'
                           )""",
                        (
                            canonical_scope,
                            f"autonomous_research:claim:{research_claim_id}",
                            f"runtime-memory-{research_claim_id}",
                        ),
                    ).lastrowid
                )
                promoted_entity_id = int(
                    conn.execute(
                        """INSERT INTO entities(
                               scope, entity_type, canonical_name, data_json
                           ) VALUES (?, 'research_claim', ?, ?)""",
                        (
                            canonical_scope,
                            f"Claim {research_claim_id}",
                            json.dumps(
                                {
                                    "statement": "Oil specification is 5W-30",
                                    "provenance": {
                                        "research_claim_id": research_claim_id
                                    },
                                },
                                ensure_ascii=False,
                            ),
                        ),
                    ).lastrowid
                )
                conn.execute(
                    """UPDATE research_claims
                       SET promoted_memory_id=?, promoted_entity_id=?,
                           promoted_at=CURRENT_TIMESTAMP
                       WHERE id=?""",
                    (
                        promoted_memory_id,
                        promoted_entity_id,
                        research_claim_id,
                    ),
                )

                graph_source = int(
                    conn.execute(
                        """INSERT INTO entities(
                               scope, entity_type, canonical_name, data_json
                           ) VALUES (?, 'vehicle', 'Runtime Car', '{}')""",
                        (canonical_scope,),
                    ).lastrowid
                )
                graph_target = int(
                    conn.execute(
                        """INSERT INTO entities(
                               scope, entity_type, canonical_name, data_json
                           ) VALUES (?, 'part', 'Runtime Filter', '{}')""",
                        (canonical_scope,),
                    ).lastrowid
                )
                conn.execute(
                    """INSERT INTO relations(
                           scope, source_entity_id, relation_type,
                           target_entity_id, confidence, evidence
                       ) VALUES (?, ?, 'uses', ?, 0.91, 'runtime graph')""",
                    (canonical_scope, graph_source, graph_target),
                )
                conn.commit()

            canonical_sync = engine.canonical_facts.sync_scope(
                scope=canonical_scope,
            )
            canonical_summary = canonical_sync.get("summary") or {}
            if canonical_summary.get("version") != "aishin-canonical-facts-v1":
                raise RuntimeError("Canonical Facts version mismatch")

            canonical_facts = engine.canonical_facts.facts(
                scope=canonical_scope,
                limit=100,
            )
            power_fact = next(
                (
                    item
                    for item in canonical_facts
                    if "engine_power:has_value" in item.get("canonical_key", "")
                ),
                None,
            )
            if (
                not power_fact
                or power_fact.get("state") != "verified"
                or power_fact.get("current_value") != "100 kW"
                or int(power_fact.get("independent_groups") or 0) != 2
            ):
                raise RuntimeError(
                    "Two independent document families must verify one "
                    "structured canonical fact"
                )

            power_groups = {
                str(evidence.get("independence_group"))
                for value in power_fact.get("values") or []
                for evidence in value.get("evidence") or []
                if evidence.get("active") and evidence.get("is_independent")
            }
            if power_groups != {
                "document_family:power-source-a",
                "document_family:power-source-b",
            }:
                raise RuntimeError(
                    "Canonical document independence groups are incorrect"
                )

            fuel_fact = next(
                (
                    item
                    for item in canonical_facts
                    if "fuel_limit:has_value" in item.get("canonical_key", "")
                ),
                None,
            )
            if (
                not fuel_fact
                or fuel_fact.get("current_value") != "130 liters"
                or int(fuel_fact.get("active_values") or 0) != 1
            ):
                raise RuntimeError(
                    "Latest document version must become current canonical value"
                )
            fuel_values = fuel_fact.get("values") or []
            old_fuel = next(
                (
                    item
                    for item in fuel_values
                    if item.get("display_value") == "120 liters"
                ),
                None,
            )
            new_fuel = next(
                (
                    item
                    for item in fuel_values
                    if item.get("display_value") == "130 liters"
                ),
                None,
            )
            if (
                not old_fuel
                or not new_fuel
                or old_fuel.get("state") != "superseded"
                or int(old_fuel.get("superseded_by_value_id") or 0)
                != int(new_fuel["id"])
            ):
                raise RuntimeError(
                    "Document version history must preserve 120 -> 130 supersession"
                )

            fuel_history = engine.canonical_facts.history(
                scope=canonical_scope,
                fact_id=int(fuel_fact["id"]),
            )
            if not any(
                item.get("event_type")
                == "value_superseded_by_document_version"
                for item in fuel_history.get("events") or []
            ):
                raise RuntimeError(
                    "Canonical history must audit document-version supersession"
                )

            interval_fact = next(
                (
                    item
                    for item in canonical_facts
                    if "service_interval:has_value"
                    in item.get("canonical_key", "")
                ),
                None,
            )
            if (
                not interval_fact
                or interval_fact.get("state") != "conflicted"
                or int(interval_fact.get("active_values") or 0) != 2
            ):
                raise RuntimeError(
                    "Independent conflicting structured values must remain conflicted"
                )

            generic_dates = [
                item
                for item in canonical_facts
                if item.get("predicate") == "has_value"
                and item.get("fact_type") == "date"
            ]
            if len(generic_dates) != 2:
                raise RuntimeError(
                    "Generic document metadata must stay family-scoped"
                )

            research_fact = next(
                (
                    item
                    for item in canonical_facts
                    if item.get("canonical_key")
                    == "research:oil_spec_claim"
                ),
                None,
            )
            if (
                not research_fact
                or research_fact.get("state") != "verified"
                or int(research_fact.get("independent_groups") or 0) != 2
            ):
                raise RuntimeError(
                    "Trusted Research evidence must project as canonical verified"
                )
            research_links = {
                (item.get("linked_type"), item.get("relation"))
                for item in research_fact.get("links") or []
            }
            if not {
                ("research_claim", "canonical_source"),
                ("memory", "derived_promotion"),
                ("graph_entity", "derived_promotion"),
            }.issubset(research_links):
                raise RuntimeError(
                    "Research -> Memory -> Graph lineage links are incomplete"
                )
            research_primary_groups = {
                str(evidence.get("independence_group"))
                for value in research_fact.get("values") or []
                for evidence in value.get("evidence") or []
                if evidence.get("active") and evidence.get("is_independent")
            }
            if research_primary_groups != {"manual:a", "manual:b"}:
                raise RuntimeError(
                    "Derived Memory/Graph promotions must not add independent votes"
                )

            graph_fact = next(
                (
                    item
                    for item in canonical_facts
                    if item.get("namespace") == "graph"
                    and item.get("predicate") == "uses"
                ),
                None,
            )
            if (
                not graph_fact
                or graph_fact.get("state") != "observed"
                or int(graph_fact.get("independent_groups") or 0) != 0
            ):
                raise RuntimeError(
                    "Knowledge Graph projection must not self-confirm as primary evidence"
                )

            canonical_claims = [
                item
                for item in engine.knowledge_lifecycle.claims(
                    scope=canonical_scope,
                    limit=100,
                )
                if item.get("origin_type") == "canonical_fact"
            ]
            if not any(
                item.get("state") == "verified"
                and "100 kW" in str(item.get("statement") or "")
                for item in canonical_claims
            ):
                raise RuntimeError(
                    "Verified canonical fact did not project into Knowledge Lifecycle"
                )

            canonical_reasoning = engine.canonical_facts.reasoning_evidence(
                "engine power 100 kW",
                scope=canonical_scope,
                limit=4,
            )
            canonical_reasoning_groups = {
                str(item.get("source_group"))
                for item in canonical_reasoning
            }
            if not {
                "document_family:power-source-a",
                "document_family:power-source-b",
            }.issubset(canonical_reasoning_groups):
                raise RuntimeError(
                    "Canonical reasoning evidence lost primary source groups"
                )
            if not all(
                item.get("trust_boundary")
                == "derived_canonical_projection"
                for item in canonical_reasoning
            ):
                raise RuntimeError(
                    "Canonical reasoning evidence must keep derived trust boundary"
                )

            canonical_api = client.get(
                "/api/assistant/canonical-facts",
                params={"scope": canonical_scope, "limit": 100},
            )
            if canonical_api.status_code != 200:
                raise RuntimeError("Canonical Facts API недоступен")
            canonical_api_data = canonical_api.json()
            if (
                canonical_api_data.get("summary", {}).get("version")
                != "aishin-canonical-facts-v1"
            ):
                raise RuntimeError("Canonical Facts API version mismatch")

            fuel_api = client.get(
                "/api/assistant/canonical-facts/history",
                params={
                    "scope": canonical_scope,
                    "fact_id": int(fuel_fact["id"]),
                },
            )
            if (
                fuel_api.status_code != 200
                or len(fuel_api.json().get("values") or []) != 2
            ):
                raise RuntimeError(
                    "Canonical fact history API did not return value lineage"
                )

            guarded_sync = client.post(
                "/api/assistant/canonical-facts/sync",
                params={"scope": canonical_scope},
            )
            if guarded_sync.status_code != 403:
                raise RuntimeError(
                    "Canonical Facts sync API must preserve local-only guard"
                )

            first_summary = engine.canonical_facts.summary(
                scope=canonical_scope,
            )
            second_sync = engine.canonical_facts.sync_scope(
                scope=canonical_scope,
            )
            second_summary = second_sync.get("summary") or {}
            for key in (
                "facts",
                "values",
                "active_evidence",
                "independent_groups",
                "superseded_values",
            ):
                if first_summary.get(key) != second_summary.get(key):
                    raise RuntimeError(
                        f"Canonical Facts sync is not idempotent for {key}"
                    )

            empty_scope = "smoke:canonical-empty"
            empty_api = client.get(
                "/api/assistant/canonical-facts",
                params={"scope": empty_scope, "limit": 20},
            )
            if (
                empty_api.status_code != 200
                or empty_api.json().get("summary", {}).get("scope")
                != empty_scope
                or int(
                    empty_api.json().get("summary", {}).get("facts") or 0
                ) != 0
            ):
                raise RuntimeError(
                    "Canonical Facts must preserve strict scope isolation"
                )

            checks["canonical_facts"] = {
                "status": "ok",
                "version": canonical_summary.get("version"),
                "verified": (
                    canonical_summary.get("states") or {}
                ).get("verified"),
                "conflicted": (
                    canonical_summary.get("states") or {}
                ).get("conflicted"),
                "superseded_values": canonical_summary.get(
                    "superseded_values"
                ),
                "cross_document_fusion": True,
                "generic_metadata_isolated": True,
                "research_memory_graph_lineage": True,
                "derived_sources_do_not_self_confirm": True,
                "history_audited": True,
                "idempotent_sync": True,
            }

            checks["response_grounding"] = {
                "status": "ok",
                "supported": supported_grounding.to_dict(),
                "numeric_mismatch": numeric_mismatch.to_dict(),
                "negation_mismatch": negation_mismatch.to_dict(),
                "no_evidence": no_evidence_grounding.to_dict(),
                "raw_tool_content_persisted": False,
            }

            proactive_conditions = client.get(
                "/api/assistant/proactive/conditions",
                params={"scope": "personal", "limit": 1},
            )
            if proactive_conditions.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/proactive/conditions returned HTTP "
                    f"{proactive_conditions.status_code}"
                )
            checks["proactive_conditions"] = {
                "status": "ok",
                "history_entries": len(proactive_conditions.json()),
            }

            graph_search = client.get(
                "/api/assistant/graph/search",
                params={
                    "scope": "relationship",
                    "query": "Айшин",
                    "limit": 5,
                },
            )
            if graph_search.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/graph/search returned HTTP "
                    f"{graph_search.status_code}"
                )
            graph_items = graph_search.json()
            if not graph_items:
                raise RuntimeError("Knowledge Graph seed Айшин не найден")
            if (
                "canonical_name" not in graph_items[0]
                or not isinstance(graph_items[0].get("data"), dict)
            ):
                raise RuntimeError(
                    "Knowledge Graph search shape несовместим с Verification"
                )
            checks["knowledge_graph_shape"] = {
                "status": "ok",
                "matches": len(graph_items),
            }

            learning_status = client.get(
                "/api/assistant/continuous-learning/status",
                params={"scope": "personal"},
            )
            if learning_status.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/continuous-learning/status returned HTTP "
                    f"{learning_status.status_code}"
                )
            learning_data = learning_status.json()
            if learning_data.get("mode") not in {
                "REALTIME", "BACKGROUND", "IDLE", "MAINTENANCE"
            }:
                raise RuntimeError("Continuous Learning вернул неизвестный mode")

            learning_cycles = client.get(
                "/api/assistant/continuous-learning/cycles",
                params={"scope": "personal", "limit": 2},
            )
            if learning_cycles.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/continuous-learning/cycles returned HTTP "
                    f"{learning_cycles.status_code}"
                )

            learning_patterns = client.get(
                "/api/assistant/continuous-learning/patterns",
                params={"scope": "personal", "limit": 2},
            )
            if learning_patterns.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/continuous-learning/patterns returned HTTP "
                    f"{learning_patterns.status_code}"
                )

            checks["continuous_learning"] = {
                "status": "ok",
                "mode": learning_data.get("mode"),
                "worker_status": learning_data.get("worker_status"),
                "cycles": len(learning_cycles.json()),
                "patterns": len(learning_patterns.json()),
            }

            learning_quality = client.get(
                "/api/assistant/continuous-learning/quality",
                params={"scope": "personal"},
            )
            if learning_quality.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/continuous-learning/quality returned HTTP "
                    f"{learning_quality.status_code}"
                )

            strategy_evolution = client.get(
                "/api/assistant/continuous-learning/strategies",
                params={"scope": "personal", "limit": 3},
            )
            if strategy_evolution.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/continuous-learning/strategies returned HTTP "
                    f"{strategy_evolution.status_code}"
                )

            checks["learning_quality"] = {
                "status": "ok",
                "quality": learning_quality.json(),
                "strategies": len(strategy_evolution.json()),
            }

            reflection_api = client.get(
                "/api/assistant/self-reflection",
                params={"scope": "personal", "limit": 2},
            )
            if reflection_api.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/self-reflection returned HTTP "
                    f"{reflection_api.status_code}"
                )

            learning_plans_api = client.get(
                "/api/assistant/learning-plans",
                params={"scope": "personal", "limit": 2},
            )
            if learning_plans_api.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/learning-plans returned HTTP "
                    f"{learning_plans_api.status_code}"
                )

            experiments_api = client.get(
                "/api/assistant/safe-experiments",
                params={"scope": "personal", "limit": 2},
            )
            if experiments_api.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/safe-experiments returned HTTP "
                    f"{experiments_api.status_code}"
                )
            for item in experiments_api.json():
                if item.get("promotion_allowed") is not False:
                    raise RuntimeError(
                        "Safe experiment must never expose automatic promotion"
                    )

            budget_api = client.get(
                "/api/assistant/context-budget",
                params={"scope": "personal", "limit": 2},
            )
            if budget_api.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/context-budget returned HTTP "
                    f"{budget_api.status_code}"
                )

            checks["reflective_learning_api"] = {
                "status": "ok",
                "reflection": bool(reflection_api.json().get("summary")),
                "plans": len(learning_plans_api.json()),
                "experiments": len(experiments_api.json()),
                "budget_reports": len(budget_api.json()),
                "automatic_promotion": False,
            }

            development_api = client.get(
                "/api/assistant/development",
                params={"scope": "personal", "days": 30},
            )
            if development_api.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/development returned HTTP "
                    f"{development_api.status_code}"
                )
            development_data = development_api.json()
            current_development = development_data.get("current") or {}
            if len(current_development.get("components") or {}) != 8:
                raise RuntimeError(
                    "Development API должен возвращать 8 направлений"
                )
            if round(
                sum((current_development.get("weights") or {}).values()), 1
            ) != 100.0:
                raise RuntimeError(
                    "Development API weights должны давать 100%"
                )
            if current_development.get("counters", {}).get(
                "confirmed_hypotheses"
            ) != 0:
                raise RuntimeError(
                    "Нельзя показывать неподтверждённые гипотезы как подтверждённые"
                )
            if "Cloud.ru не добавляет баллы" not in str(
                current_development.get("principle") or ""
            ):
                raise RuntimeError(
                    "Development API должен явно отделять развитие Айшин от Cloud.ru"
                )

            checks["development_metrics_api"] = {
                "status": "ok",
                "score": current_development.get("overall_score"),
                "components": len(current_development.get("components") or {}),
                "history": len(development_data.get("history") or []),
                "cloud_independent": True,
            }

            growth_api = client.get(
                "/api/assistant/growth",
                params={
                    "scope": "personal",
                    "skill_limit": 20,
                    "knowledge_limit": 20,
                    "specialization_limit": 20,
                    "history_limit": 20,
                },
            )
            if growth_api.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/growth returned HTTP "
                    f"{growth_api.status_code}: {growth_api.text[:300]}"
                )
            growth_data = growth_api.json()
            growth_summary = growth_data.get("summary") or {}
            if growth_summary.get("version") != "aishin-long-term-growth-v1":
                raise RuntimeError(
                    "Long Term Growth вернул несовместимую версию"
                )
            for key in ("skills", "knowledge", "specializations"):
                if not isinstance(growth_data.get(key), list):
                    raise RuntimeError(
                        f"Long Term Growth {key} должен быть списком"
                    )
            if not isinstance(growth_summary.get("principles"), list):
                raise RuntimeError(
                    "Long Term Growth должен объяснять правила честного роста"
                )
            growth_skills = growth_summary.get("skills") or {}
            growth_knowledge = growth_summary.get("knowledge") or {}
            if int(growth_skills.get("mastered") or 0) > int(
                growth_skills.get("total") or 0
            ):
                raise RuntimeError(
                    "Mastered skills не могут превышать total skills"
                )
            if int(growth_knowledge.get("trusted") or 0) > int(
                growth_knowledge.get("total") or 0
            ):
                raise RuntimeError(
                    "Trusted knowledge не может превышать total knowledge"
                )
            schema = state_data.get("state", {}).get("schema")
            checks["long_term_growth"] = {
                "status": "ok",
                "version": growth_summary.get("version"),
                "score": growth_summary.get("overall_score"),
                "skills": len(growth_data.get("skills") or []),
                "knowledge": len(growth_data.get("knowledge") or []),
                "specializations": len(
                    growth_data.get("specializations") or []
                ),
                "history": len(growth_data.get("history") or []),
            }

            intelligence_api = client.get(
                "/api/assistant/intelligence",
                params={
                    "scope": "personal",
                    "history_limit": 20,
                    "route_limit": 20,
                },
            )
            if intelligence_api.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/intelligence returned HTTP "
                    f"{intelligence_api.status_code}: "
                    f"{intelligence_api.text[:300]}"
                )
            intelligence_data = intelligence_api.json()
            current_intelligence = intelligence_data.get("current") or {}
            dimensions = current_intelligence.get("dimensions") or {}
            if current_intelligence.get("version") != (
                "aishin-cognitive-intelligence-v1"
            ):
                raise RuntimeError(
                    "Cognitive Intelligence вернул несовместимую версию"
                )
            if len(dimensions) != 8:
                raise RuntimeError(
                    "Cognitive Intelligence должен возвращать 8 способностей"
                )
            if round(
                sum(
                    float(item.get("weight") or 0.0)
                    for item in dimensions.values()
                ),
                1,
            ) != 100.0:
                raise RuntimeError(
                    "Веса 8 cognitive dimensions должны давать 100%"
                )
            for key, item in dimensions.items():
                score = float(item.get("score") or 0.0)
                if not 0.0 <= score <= 100.0:
                    raise RuntimeError(
                        f"Cognitive dimension {key} вышел за диапазон 0..100"
                    )
                if not str(item.get("why") or "").strip():
                    raise RuntimeError(
                        f"Cognitive dimension {key} не объясняет свой score"
                    )
            if not any(
                "не IQ" in str(item)
                for item in current_intelligence.get("principles") or []
            ):
                raise RuntimeError(
                    "Cognitive Intelligence должен явно отделяться от IQ"
                )

            safety_route = engine.cognitive_intelligence.route(
                request_id=engine.cognitive_traces.new_request_id(),
                scope="personal",
                query=(
                    "Проверь противоречия и перепроверь данные, "
                    "не делай вывод без подтверждения"
                ),
                intent="verification",
                base_mode="VERIFY",
                base_complexity=0.8,
                metacognition={
                    "status": "needs_verification",
                    "confidence": 0.25,
                },
            )
            if safety_route.adapted_mode != "VERIFY":
                raise RuntimeError(
                    "Adaptive Router не имеет права понижать VERIFY"
                )

            diagnose_route = engine.cognitive_intelligence.route(
                request_id=engine.cognitive_traces.new_request_id(),
                scope="personal",
                query="Диагностируй ошибку запуска системы",
                intent="action",
                base_mode="DIAGNOSE",
                base_complexity=0.8,
                metacognition={
                    "status": "cautious",
                    "confidence": 0.55,
                },
            )
            if diagnose_route.adapted_mode != "DIAGNOSE":
                raise RuntimeError(
                    "Adaptive Router не имеет права понижать DIAGNOSE"
                )

            checks["cognitive_intelligence"] = {
                "status": "ok",
                "score": current_intelligence.get("overall_score"),
                "dimensions": len(dimensions),
                "routes": len(intelligence_data.get("routes") or []),
                "transfer": len(
                    intelligence_data.get("transfer_map") or []
                ),
                "verify_preserved": True,
                "diagnose_preserved": True,
            }

            proactive_task = client.post(
                "/api/assistant/planner/tasks",
                json={
                    "title": "Runtime Smoke — просроченная задача",
                    "description": "Проверка Proactive Intelligence lifecycle",
                    "scope": "personal",
                    "priority": 0.9,
                    "due_at": "2020-01-01T00:00:00+00:00",
                },
            )
            if proactive_task.status_code != 200:
                raise RuntimeError(
                    "Не удалось создать deterministic overdue task для "
                    "Proactive Intelligence smoke"
                )
            proactive_task_id = int(proactive_task.json()["id"])

            proactive_scan_data = engine.proactive_intelligence.evaluate(
                scope="personal",
                trigger="runtime_smoke",
            ).to_dict()
            if not 0.0 <= float(
                proactive_scan_data.get("awareness_score") or 0.0
            ) <= 100.0:
                raise RuntimeError(
                    "Situation Awareness должен оставаться в диапазоне 0..100"
                )

            proactive_api = client.get(
                "/api/assistant/proactive-intelligence",
                params={
                    "scope": "personal",
                    "incident_limit": 100,
                    "signal_limit": 100,
                    "run_limit": 30,
                },
            )
            if proactive_api.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/proactive-intelligence returned HTTP "
                    f"{proactive_api.status_code}: {proactive_api.text[:300]}"
                )
            proactive_data = proactive_api.json()
            if proactive_data.get("version") != (
                "aishin-proactive-intelligence-v1"
            ):
                raise RuntimeError(
                    "Proactive Intelligence вернул несовместимую версию"
                )
            summary = proactive_data.get("summary") or {}
            threshold = float(summary.get("attention_threshold") or 0.0)
            if not 0.36 <= threshold <= 0.70:
                raise RuntimeError(
                    "Attention threshold вышел за безопасный диапазон"
                )
            if not isinstance(proactive_data.get("incidents"), list):
                raise RuntimeError(
                    "Proactive Intelligence incidents должен быть списком"
                )
            if not isinstance(proactive_data.get("expectations"), list):
                raise RuntimeError(
                    "Proactive Intelligence expectations должен быть списком"
                )
            if not isinstance(proactive_data.get("signals"), list):
                raise RuntimeError(
                    "Proactive Intelligence signals должен быть списком"
                )

            active_incidents = [
                item
                for item in proactive_data.get("incidents") or []
                if item.get("status") == "active"
            ]
            overdue_incident = next(
                (
                    item
                    for item in active_incidents
                    if item.get("subject_type") == "task"
                    and str(item.get("subject_id")) == str(proactive_task_id)
                    and item.get("incident_type") == "task_overdue"
                ),
                None,
            )
            if overdue_incident is None:
                raise RuntimeError(
                    "Proactive Intelligence не обнаружил deterministic overdue task"
                )

            for key in (
                "severity",
                "confidence",
                "impact",
                "urgency",
                "risk_score",
                "attention_score",
            ):
                value = float(overdue_incident.get(key) or 0.0)
                if not 0.0 <= value <= 1.0:
                    raise RuntimeError(
                        f"Proactive incident {key} вышел за диапазон 0..1"
                    )
            if not isinstance(overdue_incident.get("evidence"), list):
                raise RuntimeError(
                    "Proactive incident должен хранить evidence"
                )
            decision_id = overdue_incident.get("decision_id")
            if decision_id:
                decision = next(
                    (
                        item
                        for item in engine.proactive.history(
                            scope="personal",
                            limit=300,
                        )
                        if int(item["id"]) == int(decision_id)
                    ),
                    None,
                )
                if decision and decision.get("tool_name"):
                    raise RuntimeError(
                        "00.00.07 не должен автоматически создавать "
                        "исполняемый tool proposal"
                    )

            proactive_feedback = engine.proactive_intelligence.feedback(
                int(overdue_incident["id"]),
                scope="personal",
                feedback="noisy",
                reason="runtime smoke suppression check",
            )
            feedback_incident = proactive_feedback.get("incident") or {}
            if feedback_incident.get("status") != "snoozed":
                raise RuntimeError(
                    "Noisy feedback должен временно подавлять incident"
                )
            if not feedback_incident.get("snoozed_until"):
                raise RuntimeError(
                    "Noisy feedback должен сохранять snoozed_until"
                )

            engine.proactive_intelligence.evaluate(
                scope="personal",
                trigger="runtime_smoke_rescan",
            )
            incident_after_rescan = engine.proactive_intelligence.incident(
                int(overdue_incident["id"]),
                scope="personal",
            ) or {}
            if incident_after_rescan.get("status") != "snoozed":
                raise RuntimeError(
                    "Suppressed incident не должен немедленно возвращаться active"
                )

            close_proactive_task = client.post(
                f"/api/assistant/planner/tasks/{proactive_task_id}/status",
                json={
                    "scope": "personal",
                    "status": "completed",
                    "blocked_reason": "",
                },
            )
            if close_proactive_task.status_code != 200:
                raise RuntimeError(
                    "Не удалось завершить proactive smoke task"
                )
            engine.proactive_intelligence.evaluate(
                scope="personal",
                trigger="runtime_smoke_cleanup",
            )

            checks["proactive_intelligence"] = {
                "status": "ok",
                "version": proactive_data.get("version"),
                "awareness": summary.get("awareness_score"),
                "incidents": len(proactive_data.get("incidents") or []),
                "signals": len(proactive_data.get("signals") or []),
                "expectations": len(
                    proactive_data.get("expectations") or []
                ),
                "threshold": threshold,
                "feedback_suppression": True,
                "auto_execution": False,
            }

            evolution_scope = "__evolution_smoke__"
            with connect() as conn:
                for table in (
                    "evolution_assignments",
                    "evolution_variants",
                    "evolution_curriculum",
                    "evolution_transfers",
                    "evolution_cycles",
                    "evolution_events",
                    "evolution_capabilities",
                    "evolution_state",
                    "cognitive_intelligence_routes",
                ):
                    conn.execute(
                        f"DELETE FROM {table} WHERE scope=?",
                        (evolution_scope,),
                    )
                conn.execute(
                    """DELETE FROM learning_plans
                       WHERE scope=? AND target_metric LIKE 'evolution:%'""",
                    (evolution_scope,),
                )
                for index in range(12):
                    outcome = 0.54 + (index % 3) * 0.015
                    conn.execute(
                        """INSERT INTO cognitive_intelligence_routes(
                               request_id, scope, task_family, intent,
                               base_mode, adapted_mode, outcome_score,
                               successful, unresolved_count, completed_at
                           ) VALUES (?, ?, 'software', 'analysis',
                                     'FAST', 'FAST', ?, 0, 0,
                                     CURRENT_TIMESTAMP)""",
                        (
                            f"evolution-seed-{index}",
                            evolution_scope,
                            outcome,
                        ),
                    )
                conn.commit()

            first_evolution_cycle = engine.evolution.run_cycle(
                scope=evolution_scope,
                trigger="runtime_smoke_seed",
            )
            if not 0.0 <= float(first_evolution_cycle.evolution_score) <= 100.0:
                raise RuntimeError(
                    "Evolution score должен оставаться в диапазоне 0..100"
                )
            if not 0.0 <= float(first_evolution_cycle.stability_score) <= 100.0:
                raise RuntimeError(
                    "Evolution stability должен оставаться в диапазоне 0..100"
                )
            if not 0.0 <= float(first_evolution_cycle.plasticity_score) <= 100.0:
                raise RuntimeError(
                    "Evolution plasticity должен оставаться в диапазоне 0..100"
                )

            evolution_api = client.get(
                "/api/assistant/evolution",
                params={
                    "scope": evolution_scope,
                    "capability_limit": 40,
                    "variant_limit": 80,
                    "curriculum_limit": 80,
                    "transfer_limit": 80,
                    "cycle_limit": 80,
                },
            )
            if evolution_api.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/evolution returned HTTP "
                    f"{evolution_api.status_code}: "
                    f"{evolution_api.text[:300]}"
                )
            evolution_data = evolution_api.json()
            if evolution_data.get("version") != "aishin-evolution-engine-v1":
                raise RuntimeError(
                    "Evolution Engine вернул несовместимую версию"
                )
            if evolution_data.get("formula_version") != "bounded-meta-learning-v1":
                raise RuntimeError(
                    "Evolution Engine вернул несовместимую формулу"
                )
            if not isinstance(evolution_data.get("capabilities"), list):
                raise RuntimeError(
                    "Evolution capabilities должен быть списком"
                )
            if not isinstance(evolution_data.get("curriculum"), list):
                raise RuntimeError(
                    "Evolution curriculum должен быть списком"
                )
            if not any(
                item.get("family") == "software"
                and int(item.get("sample_count") or 0) >= 12
                for item in evolution_data.get("capabilities") or []
            ):
                raise RuntimeError(
                    "Evolution Engine не построил capability из route outcomes"
                )

            challengers = [
                item
                for item in evolution_data.get("variants") or []
                if item.get("family") == "software"
                and item.get("lifecycle") == "challenger"
            ]
            if not challengers:
                raise RuntimeError(
                    "Evolution Engine не создал challenger для слабой family"
                )
            challenger = challengers[0]
            challenger_id = int(challenger["id"])
            baseline = float(challenger.get("baseline_fitness") or 0.0)

            verify_hint = engine.evolution.routing_hint(
                scope=evolution_scope,
                family="software",
                base_mode="VERIFY",
                request_id="evolution-safety-verify",
                base_complexity=0.9,
            )
            if verify_hint.get("preferred_mode") != "VERIFY":
                raise RuntimeError(
                    "Evolution Engine не имеет права понижать VERIFY"
                )
            diagnose_hint = engine.evolution.routing_hint(
                scope=evolution_scope,
                family="software",
                base_mode="DIAGNOSE",
                request_id="evolution-safety-diagnose",
                base_complexity=0.9,
            )
            if diagnose_hint.get("preferred_mode") != "DIAGNOSE":
                raise RuntimeError(
                    "Evolution Engine не имеет права понижать DIAGNOSE"
                )
            for hint in (verify_hint, diagnose_hint):
                multiplier = float(hint.get("context_multiplier") or 1.0)
                if not 0.90 <= multiplier <= 1.20:
                    raise RuntimeError(
                        "Evolution context policy вышла за safe bounds"
                    )

            with connect() as conn:
                conn.execute(
                    """DELETE FROM evolution_assignments
                       WHERE scope=? AND variant_id=?""",
                    (evolution_scope, challenger_id),
                )
                high_outcome = min(0.98, baseline + 0.18)
                for index in range(10):
                    conn.execute(
                        """INSERT INTO evolution_assignments(
                               request_id, scope, family, variant_id,
                               assignment_type, policy_json,
                               baseline_fitness, outcome_score, successful,
                               unresolved_count, completed_at
                           ) VALUES (?, ?, 'software', ?, 'challenger',
                                     '{}', ?, ?, 1, 0,
                                     CURRENT_TIMESTAMP)""",
                        (
                            f"evolution-promote-{index}",
                            evolution_scope,
                            challenger_id,
                            baseline,
                            high_outcome,
                        ),
                    )
                conn.execute(
                    """UPDATE evolution_variants
                       SET evidence_count=10, wins=10, losses=0,
                           unresolved_total=0, observed_fitness=?
                       WHERE id=?""",
                    (high_outcome, challenger_id),
                )
                conn.commit()

            promote_cycle = engine.evolution.run_cycle(
                scope=evolution_scope,
                trigger="runtime_smoke_promote",
            )
            promoted = engine.evolution.variants(
                scope=evolution_scope,
                lifecycle="champion",
                limit=10,
            )
            champion = next(
                (
                    item for item in promoted
                    if int(item["id"]) == challenger_id
                ),
                None,
            )
            if champion is None or promote_cycle.variants_promoted < 1:
                raise RuntimeError(
                    "Conservative Evolution evidence gate не promoted challenger"
                )

            with connect() as conn:
                for index in range(5):
                    conn.execute(
                        """INSERT INTO evolution_assignments(
                               request_id, scope, family, variant_id,
                               assignment_type, policy_json,
                               baseline_fitness, outcome_score, successful,
                               unresolved_count, completed_at
                           ) VALUES (?, ?, 'software', ?, 'champion',
                                     '{}', ?, 0.28, 0, 1,
                                     CURRENT_TIMESTAMP)""",
                        (
                            f"evolution-regression-{index}",
                            evolution_scope,
                            challenger_id,
                            baseline,
                        ),
                    )
                conn.execute(
                    """UPDATE evolution_variants
                       SET evidence_count=15, losses=losses+5,
                           unresolved_total=unresolved_total+5
                       WHERE id=?""",
                    (challenger_id,),
                )
                conn.commit()

            rollback_cycle = engine.evolution.run_cycle(
                scope=evolution_scope,
                trigger="runtime_smoke_rollback",
            )
            after_rollback = next(
                (
                    item
                    for item in engine.evolution.variants(
                        scope=evolution_scope,
                        limit=100,
                    )
                    if int(item["id"]) == challenger_id
                ),
                {},
            )
            if (
                after_rollback.get("lifecycle") != "rolled_back"
                or rollback_cycle.variants_rolled_back < 1
            ):
                raise RuntimeError(
                    "Evolution Engine не откатил деградировавший champion"
                )

            evolution_state = engine.evolution.state(scope=evolution_scope)
            if int(evolution_state.get("generation") or 0) < 3:
                raise RuntimeError(
                    "Evolution generation не изменилась после promotion/rollback"
                )
            principles = evolution_data.get("principles") or []
            if not any("исходный код" in str(item) for item in principles):
                raise RuntimeError(
                    "Evolution должен явно запрещать self-modifying source code"
                )

            checks["evolution_engine"] = {
                "status": "ok",
                "version": evolution_data.get("version"),
                "generation": evolution_state.get("generation"),
                "score": evolution_state.get("evolution_score"),
                "capabilities": len(
                    evolution_data.get("capabilities") or []
                ),
                "curriculum": len(
                    evolution_data.get("curriculum") or []
                ),
                "challenger_created": True,
                "conservative_promotion": True,
                "verify_preserved": True,
                "diagnose_preserved": True,
                "rollback_verified": True,
                "self_modifying_code": False,
            }

            research_scope = "__research_smoke__"
            with connect() as conn:
                for table in (
                    "research_contradictions",
                    "research_claims",
                    "research_evidence",
                    "research_sessions",
                    "research_gaps",
                    "research_cycles",
                    "research_events",
                    "research_sources",
                    "research_state",
                ):
                    conn.execute(
                        f"DELETE FROM {table} WHERE scope=?",
                        (research_scope,),
                    )
                conn.execute(
                    "DELETE FROM memories WHERE scope=?",
                    (research_scope,),
                )
                conn.execute(
                    "DELETE FROM entities WHERE scope=?",
                    (research_scope,),
                )
                conn.commit()

            research_question = (
                "Контрольный факт: система Alpha использует "
                "трёхисточниковую проверку evidence."
            )
            for index, group in enumerate(("source-a", "source-b", "source-c"), start=1):
                engine.research.register_source(
                    scope=research_scope,
                    source_key=f"manual:{index}",
                    source_type="manual_reference",
                    label=f"Independent source {index}",
                    independent_group=group,
                    trust_prior=0.99,
                    enabled=True,
                    auto_read=True,
                    metadata={
                        "content": (
                            research_question
                            + f" Независимое подтверждение #{index}."
                        )
                    },
                )

            research_run = engine.research.research_query(
                scope=research_scope,
                question=research_question,
                trigger="runtime_smoke",
                synthesize=False,
            )
            if research_run.evidence_count < 3:
                raise RuntimeError(
                    "Research Engine не собрал 3 независимых evidence"
                )
            research_evidence = engine.research.evidence(
                scope=research_scope,
                session_id=research_run.session_id,
                limit=20,
            )
            evidence_groups = {
                item.get("source_group")
                for item in research_evidence
            }
            if not {"source-a", "source-b", "source-c"}.issubset(
                evidence_groups
            ):
                raise RuntimeError(
                    "Research Engine потерял независимость source groups"
                )

            support_ids = [
                int(item["id"])
                for item in research_evidence
                if item.get("source_group") in {
                    "source-a",
                    "source-b",
                    "source-c",
                }
            ][:3]
            trusted_claim = engine.research.evaluate_claim(
                scope=research_scope,
                session_id=research_run.session_id,
                statement=(
                    "Система Alpha использует трёхисточниковую "
                    "проверку evidence."
                ),
                support_evidence_ids=support_ids,
                contradiction_evidence_ids=[],
                missing=[],
            )
            if trusted_claim.get("status") != "trusted":
                raise RuntimeError(
                    "Evidence Gate не promoted claim с 3 сильными "
                    "независимыми источниками"
                )
            promoted_memory_id = trusted_claim.get(
                "promoted_memory_id"
            )
            if not promoted_memory_id:
                raise RuntimeError(
                    "Trusted research claim не получил provenance memory"
                )
            with connect() as conn:
                promoted_memory = conn.execute(
                    """SELECT * FROM memories
                       WHERE id=? AND scope=? AND status='active'""",
                    (int(promoted_memory_id), research_scope),
                ).fetchone()
            if promoted_memory is None:
                raise RuntimeError(
                    "Promoted research memory не найдена"
                )

            counter_evidence = engine.research.add_evidence(
                scope=research_scope,
                session_id=research_run.session_id,
                source_type="manual_reference",
                source_ref="manual:counter",
                source_group="source-counter",
                title="Independent counter source",
                content=(
                    "Контрольный факт: система Alpha НЕ использует "
                    "трёхисточниковую проверку evidence."
                ),
                reliability=0.99,
                relevance=1.0,
                freshness=1.0,
                independence=1.0,
                metadata={"runtime_smoke": True},
            )
            conflicted_claim = engine.research.evaluate_claim(
                scope=research_scope,
                session_id=research_run.session_id,
                statement=(
                    "Система Alpha использует трёхисточниковую "
                    "проверку evidence."
                ),
                support_evidence_ids=support_ids,
                contradiction_evidence_ids=[
                    int(counter_evidence["id"])
                ],
                missing=[],
            )
            if conflicted_claim.get("status") != "conflicted":
                raise RuntimeError(
                    "Counter-evidence должен переводить claim в conflicted"
                )
            if not engine.research.contradictions(
                scope=research_scope,
                session_id=research_run.session_id,
                status="open",
                limit=20,
            ):
                raise RuntimeError(
                    "Contradiction Matrix не сохранила counter-evidence"
                )
            with connect() as conn:
                memory_after_conflict = conn.execute(
                    """SELECT status FROM memories
                       WHERE id=? AND scope=?""",
                    (int(promoted_memory_id), research_scope),
                ).fetchone()
            if (
                memory_after_conflict is None
                or memory_after_conflict["status"] != "active"
            ):
                raise RuntimeError(
                    "Research regression не должна молча удалять "
                    "ранее promoted memory"
                )

            network_source = engine.research.register_source(
                scope=research_scope,
                source_key="external:disabled-adapter",
                source_type="external_connector",
                label="External adapter placeholder",
                independent_group="external",
                trust_prior=0.9,
                enabled=True,
                auto_read=True,
                metadata={"url": "https://example.invalid/"},
            )
            if network_source.get("source_type") != "external_connector":
                raise RuntimeError(
                    "Research source registry shape несовместим"
                )
            research_dashboard = engine.research.dashboard(
                scope=research_scope,
                gap_limit=40,
                session_limit=40,
                claim_limit=40,
                evidence_limit=60,
                cycle_limit=20,
            )
            research_summary = research_dashboard.get("summary") or {}
            research_score = float(
                research_summary.get("research_score") or 0.0
            )
            if not 0.0 <= research_score <= 100.0:
                raise RuntimeError(
                    "Research Score должен оставаться в диапазоне 0..100"
                )
            if not any(
                "Model synthesis is never evidence" in str(item)
                for item in research_dashboard.get("principles") or []
            ):
                raise RuntimeError(
                    "Research Engine должен явно запрещать "
                    "self-confirmation модели"
                )
            research_api = client.get(
                "/api/assistant/research",
                params={"scope": research_scope},
            )
            if research_api.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/research returned HTTP "
                    f"{research_api.status_code}"
                )
            if research_api.json().get("version") != (
                "aishin-autonomous-research-v1"
            ):
                raise RuntimeError(
                    "Research API вернул несовместимую версию"
                )

            checks["autonomous_research"] = {
                "status": "ok",
                "version": research_dashboard.get("version"),
                "research_score": research_score,
                "evidence": len(research_evidence),
                "trusted_promotion": True,
                "counter_evidence_blocks_trust": True,
                "promoted_memory_preserved_for_review": True,
                "external_network_used": False,
            }

            communication_scope = "smoke:communication-v1"
            with connect() as conn:
                conn.execute(
                    """DELETE FROM communication_feedback
                       WHERE scope=?""",
                    (communication_scope,),
                )
                conn.execute(
                    """DELETE FROM communication_events
                       WHERE scope=?""",
                    (communication_scope,),
                )
                conn.execute(
                    """DELETE FROM communication_turns
                       WHERE scope=?""",
                    (communication_scope,),
                )
                conn.execute(
                    """DELETE FROM communication_preferences
                       WHERE scope=?""",
                    (communication_scope,),
                )
                conn.execute(
                    """DELETE FROM communication_skills
                       WHERE scope=?""",
                    (communication_scope,),
                )
                conn.execute(
                    """DELETE FROM communication_state
                       WHERE scope=?""",
                    (communication_scope,),
                )
                conn.commit()

            communication_boot = engine.communication.bootstrap(
                scope=communication_scope,
            )
            if float(
                communication_boot.get("communication_score") or 0.0
            ) != 0.0:
                raise RuntimeError(
                    "Communication Score без outcome evidence должен быть 0"
                )

            communication_plan_1 = engine.communication.prepare(
                scope=communication_scope,
                request_id="smoke-communication-1",
                user_message=(
                    "Объясни профессионально максимально подробно, "
                    "как работает этот модуль."
                ),
                base_intent="question",
                logic_mode="DEEP",
            )
            if communication_plan_1.depth != "expert":
                raise RuntimeError(
                    "Expert-маркеры должны выбирать expert depth"
                )
            engine.communication.record_response(
                scope=communication_scope,
                request_id="smoke-communication-1",
                assistant_message=(
                    "Сначала дам вывод, затем архитектуру, ограничения "
                    "и проверяемые последствия."
                ),
            )

            communication_plan_2 = engine.communication.prepare(
                scope=communication_scope,
                request_id="smoke-communication-2",
                user_message="Не понял, непонятно. Объясни проще на примере.",
                base_intent="question",
                logic_mode="FAST",
            )
            if communication_plan_2.user_signals.get(
                "positive_signal", {}
            ).get("active"):
                raise RuntimeError(
                    "Негативное «непонятно» не должно считаться "
                    "positive_signal"
                )
            if communication_plan_2.strategy != "clarification_recovery":
                raise RuntimeError(
                    "После «не понял» должен включаться "
                    "clarification_recovery"
                )
            if communication_plan_2.explanation_style != (
                "simple_after_clarification"
            ):
                raise RuntimeError(
                    "Clarification recovery должен менять explanation style"
                )
            engine.communication.record_response(
                scope=communication_scope,
                request_id="smoke-communication-2",
                assistant_message=(
                    "Простой пример: сообщение — это вход, стратегия — "
                    "способ объяснения, а feedback показывает, сработал ли он."
                ),
            )

            communication_turns = engine.communication.turns(
                scope=communication_scope,
                limit=10,
            )
            first_turn = next(
                item for item in communication_turns
                if item["request_id"] == "smoke-communication-1"
            )
            if first_turn.get("outcome") != "clarification_needed":
                raise RuntimeError(
                    "Следующая реплика «не понял» должна оценить "
                    "предыдущий turn как clarification_needed"
                )

            feedback_1 = engine.communication.feedback(
                int(first_turn["id"]),
                scope=communication_scope,
                feedback="useful",
                reason="smoke override",
            )
            if feedback_1.get("turn", {}).get("outcome") != "useful":
                raise RuntimeError(
                    "Communication feedback engine не применил useful feedback"
                )
            skills_after_first_feedback = engine.communication.skills(
                scope=communication_scope,
                limit=30,
            )
            complex_skill_1 = next(
                item for item in skills_after_first_feedback
                if item["skill_key"] == "complex_explanation"
            )
            sample_count_1 = int(complex_skill_1["sample_count"] or 0)

            feedback_2 = engine.communication.feedback(
                int(first_turn["id"]),
                scope=communication_scope,
                feedback="wrong",
                reason="smoke override second time",
            )
            if feedback_2.get("turn", {}).get("outcome") != (
                "correction_needed"
            ):
                raise RuntimeError(
                    "Communication feedback engine не применил override"
                )
            skills_after_second_feedback = engine.communication.skills(
                scope=communication_scope,
                limit=30,
            )
            complex_skill_2 = next(
                item for item in skills_after_second_feedback
                if item["skill_key"] == "complex_explanation"
            )
            if int(complex_skill_2["sample_count"] or 0) != sample_count_1:
                raise RuntimeError(
                    "Повторный feedback одного turn не должен "
                    "удваивать skill sample_count"
                )

            communication_api = client.get(
                "/api/assistant/communication",
                params={"scope": communication_scope},
            )
            if communication_api.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/communication returned HTTP "
                    f"{communication_api.status_code}"
                )
            communication_data = communication_api.json()
            if communication_data.get("version") != (
                "aishin-communication-intelligence-v1"
            ):
                raise RuntimeError(
                    "Communication API вернул несовместимую версию"
                )
            if communication_data.get("persona", {}).get(
                "phrase_library_version"
            ) != "2.2.0":
                raise RuntimeError(
                    "Communication Runtime должен загрузить "
                    "AISHIN_PHRASE_LIBRARY_RU 2.2.0"
                )
            if not isinstance(
                communication_data.get("skills"), list
            ):
                raise RuntimeError(
                    "Communication API должен возвращать skills"
                )

            checks["communication_intelligence"] = {
                "status": "ok",
                "version": communication_data.get("version"),
                "score": communication_data.get(
                    "summary", {}
                ).get("communication_score"),
                "expert_depth": True,
                "clarification_recovery": True,
                "negative_signal_disambiguated": True,
                "feedback_override_no_double_count": True,
                "phrase_library_version": "2.2.0",
            }

            document_scope = "smoke:documents-v1"
            with connect() as conn:
                conn.execute(
                    "DELETE FROM research_contradictions WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM research_claims WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM research_evidence WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM research_sessions WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM research_gaps WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM research_cycles WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM research_events WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM research_sources WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM research_state WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    """DELETE FROM document_chunk_vectors
                       WHERE scope=?""",
                    (document_scope,),
                )
                conn.execute(
                    """DELETE FROM document_contradictions
                       WHERE scope=?""",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM document_facts WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM document_chunks WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM document_sections WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM document_pages WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    """DELETE FROM document_ingestion_runs
                       WHERE scope=?""",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM document_events WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM documents WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM document_state WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM relations WHERE scope=?",
                    (document_scope,),
                )
                conn.execute(
                    "DELETE FROM entities WHERE scope=?",
                    (document_scope,),
                )
                conn.commit()

            document_body = (
                "РЕГЛАМЕНТ ГСМ\n"
                "Норма расхода: 12 литров на 100 км\n"
                "Ответственный: Отдел эксплуатации\n"
                "Дата утверждения: 01.01.2025\n"
                "Настоящий регламент определяет порядок контроля "
                "расхода топлива и проверки путевых листов. "
                "Каждое значение должно сверяться с первичным "
                "документом и карточкой автомобиля.\n"
                + "\n".join(
                    f"Контрольный пункт {i}: проверка пробега, топлива "
                    "и подтверждающих документов выполняется до "
                    "закрытия отчётного периода."
                    for i in range(1, 18)
                )
            ).encode("utf-8")
            document_v1 = engine.documents.ingest_bytes(
                scope=document_scope,
                filename="Регламент ГСМ 2025.txt",
                data=document_body,
                media_type="text/plain",
                trigger="runtime_smoke",
                enrich_with_ai=False,
                build_semantic_index=False,
            )
            if document_v1.status != "studied":
                raise RuntimeError(
                    "Качественный TXT должен пройти Document Quality Gate"
                )
            if document_v1.quality_score < 0.70:
                raise RuntimeError(
                    "Document quality unexpectedly below gate"
                )
            detail_v1 = engine.documents.document_detail(
                document_v1.document_id,
                scope=document_scope,
            )
            chunks_v1 = detail_v1.get("chunks") or []
            if not chunks_v1:
                raise RuntimeError(
                    "Document Intelligence не создал chunks"
                )
            if not all(
                isinstance(item.get("provenance"), dict)
                and item.get("provenance", {}).get("document_id")
                == document_v1.document_id
                for item in chunks_v1
            ):
                raise RuntimeError(
                    "Каждый document chunk должен иметь provenance"
                )
            if not any(
                item.get("status") == "grounded"
                and item.get("fact_type") == "key_value"
                for item in detail_v1.get("facts") or []
            ):
                raise RuntimeError(
                    "Quality Gate не создал grounded key/value facts"
                )

            exact_duplicate = engine.documents.ingest_bytes(
                scope=document_scope,
                filename="Регламент ГСМ 2025 copy.txt",
                data=document_body,
                media_type="text/plain",
                trigger="runtime_smoke_duplicate",
                enrich_with_ai=False,
                build_semantic_index=False,
            )
            if (
                not exact_duplicate.duplicate
                or exact_duplicate.document_id
                != document_v1.document_id
            ):
                raise RuntimeError(
                    "SHA-256 duplicate должен переиспользовать document_id"
                )

            document_body_v2 = document_body.decode("utf-8").replace(
                "12 литров на 100 км",
                "13 литров на 100 км",
            ).replace(
                "01.01.2025",
                "01.01.2026",
            ).encode("utf-8")
            document_v2 = engine.documents.ingest_bytes(
                scope=document_scope,
                filename="Регламент ГСМ 2026.txt",
                data=document_body_v2,
                media_type="text/plain",
                trigger="runtime_smoke_version",
                enrich_with_ai=False,
                build_semantic_index=False,
            )
            detail_v2 = engine.documents.document_detail(
                document_v2.document_id,
                scope=document_scope,
            )
            if (
                detail_v2.get("document", {}).get("previous_version_id")
                != document_v1.document_id
            ):
                raise RuntimeError(
                    "Version Intelligence не связал новую редакцию "
                    "с предыдущей"
                )
            document_conflicts = engine.documents.contradictions(
                scope=document_scope,
                status="open",
                limit=50,
            )
            if not document_conflicts:
                raise RuntimeError(
                    "Разные grounded значения двух версий должны "
                    "создать document contradiction"
                )

            document_search = engine.documents.search(
                "норма расхода топлива",
                scope=document_scope,
                limit=10,
            )
            if not document_search:
                raise RuntimeError(
                    "Lexical document retrieval не нашёл изученный регламент"
                )
            if not isinstance(
                document_search[0].get("provenance"),
                dict,
            ):
                raise RuntimeError(
                    "Document retrieval должен сохранять provenance"
                )

            document_research = engine.research.research_query(
                scope=document_scope,
                question="норма расхода топлива",
                trigger="runtime_smoke_document",
                synthesize=False,
                offline_only=True,
            )
            document_evidence = engine.research.evidence(
                scope=document_scope,
                session_id=document_research.session_id,
                limit=50,
            )
            if not any(
                item.get("source_type") == "document"
                and isinstance(
                    item.get("metadata", {}).get("provenance"),
                    dict,
                )
                for item in document_evidence
            ):
                raise RuntimeError(
                    "Research Intelligence не получил document evidence "
                    "с provenance"
                )

            docx_buffer = BytesIO()
            docx = SmokeDocxDocument()
            docx.add_heading("ТЕХНИЧЕСКОЕ РУКОВОДСТВО", level=1)
            docx.add_paragraph("Контрольный параметр: 42 единицы")
            for i in range(1, 16):
                docx.add_paragraph(
                    f"Раздел {i}. Проверяемый текст руководства "
                    "содержит описание процедуры, условия применения "
                    "и ссылку на первичный документ."
                )
            docx.save(docx_buffer)
            docx_result = engine.documents.ingest_bytes(
                scope=document_scope,
                filename="Техническое руководство 2025.docx",
                data=docx_buffer.getvalue(),
                media_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                trigger="runtime_smoke_docx",
                enrich_with_ai=False,
                build_semantic_index=False,
            )
            if (
                docx_result.status != "studied"
                or docx_result.section_count < 1
                or docx_result.chunk_count < 1
            ):
                raise RuntimeError(
                    "DOCX parser должен сохранить структуру и chunks"
                )

            xlsx_buffer = BytesIO()
            workbook = SmokeWorkbook()
            sheet = workbook.active
            sheet.title = "Нормативы"
            sheet.append(["Параметр", "Значение", "Комментарий"])
            for i in range(1, 25):
                sheet.append([
                    f"Норма {i}",
                    i * 10,
                    (
                        "Проверяемое значение для структурного "
                        "извлечения таблицы"
                    ),
                ])
            workbook.save(xlsx_buffer)
            xlsx_result = engine.documents.ingest_bytes(
                scope=document_scope,
                filename="Нормативы 2025.xlsx",
                data=xlsx_buffer.getvalue(),
                media_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                trigger="runtime_smoke_xlsx",
                enrich_with_ai=False,
                build_semantic_index=False,
            )
            if (
                xlsx_result.status != "studied"
                or xlsx_result.page_count != 1
                or xlsx_result.section_count != 1
            ):
                raise RuntimeError(
                    "XLSX parser должен сохранять worksheet как section"
                )

            pdf_buffer = BytesIO()
            pdf_writer = PdfWriter()
            pdf_writer.add_blank_page(width=300, height=400)
            pdf_writer.write(pdf_buffer)
            pdf_result = engine.documents.ingest_bytes(
                scope=document_scope,
                filename="Скан PDF без текста.pdf",
                data=pdf_buffer.getvalue(),
                media_type="application/pdf",
                trigger="runtime_smoke_pdf_ocr",
                enrich_with_ai=False,
                build_semantic_index=False,
            )
            if (
                pdf_result.status != "needs_ocr"
                or not pdf_result.ocr_required
            ):
                raise RuntimeError(
                    "PDF без текстового слоя должен получить needs_ocr"
                )

            png_1x1 = base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwC"
                "AAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
            )
            scan_result = engine.documents.ingest_bytes(
                scope=document_scope,
                filename="Скан без OCR.png",
                data=png_1x1,
                media_type="image/png",
                trigger="runtime_smoke_ocr",
                enrich_with_ai=False,
                build_semantic_index=False,
            )
            if (
                scan_result.status != "needs_ocr"
                or not scan_result.ocr_required
                or scan_result.quality_score > 0.28
            ):
                raise RuntimeError(
                    "Image без OCR adapter должен оставаться needs_ocr"
                )
            scan_detail = engine.documents.document_detail(
                scan_result.document_id,
                scope=document_scope,
            )
            if any(
                item.get("status") == "grounded"
                for item in scan_detail.get("facts") or []
            ):
                raise RuntimeError(
                    "OCR-required image не должен создавать grounded facts"
                )

            documents_api = client.get(
                "/api/assistant/documents",
                params={"scope": document_scope},
            )
            if documents_api.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/documents returned HTTP "
                    f"{documents_api.status_code}"
                )
            document_dashboard = documents_api.json()
            if document_dashboard.get("version") != (
                "aishin-document-intelligence-v1"
            ):
                raise RuntimeError(
                    "Document Intelligence API вернул несовместимую версию"
                )
            document_detail_api = client.get(
                f"/api/assistant/documents/detail/"
                f"{document_v1.document_id}",
                params={"scope": document_scope},
            )
            if document_detail_api.status_code != 200:
                raise RuntimeError(
                    "Document detail API недоступен"
                )
            document_search_api = client.get(
                "/api/assistant/documents/search",
                params={
                    "scope": document_scope,
                    "query": "норма расхода топлива",
                    "limit": 10,
                },
            )
            if (
                document_search_api.status_code != 200
                or not document_search_api.json()
            ):
                raise RuntimeError(
                    "Document search API не вернул grounded retrieval"
                )

            reasoning_docs = engine.documents.reasoning_evidence(
                "норма расхода топлива",
                scope=document_scope,
                limit=6,
            )
            reasoning_evidence = reasoning_docs.get("evidence") or []
            if not reasoning_evidence:
                raise RuntimeError(
                    "Document evidence не входит в ранний reasoning controller"
                )
            if not all(
                item.get("source") == "document"
                and item.get("trust_boundary")
                == "untrusted_document_data"
                and isinstance(item.get("provenance"), dict)
                for item in reasoning_evidence
            ):
                raise RuntimeError(
                    "Document reasoning evidence потерял trust boundary/provenance"
                )

            logic_plan_docs = engine.logic.prepare(
                "проверить норму расхода топлива",
                intent="verification",
                metacognition_status="cautious",
                metacognition_confidence=0.55,
                contradiction_count=0,
                planner_notices=[],
            )
            logic_docs = engine.logic.finalize(
                scope=document_scope,
                query="проверить норму расхода топлива",
                intent="verification",
                plan=logic_plan_docs,
                memories=[],
                final_metacognition={
                    "status": "cautious",
                    "confidence": 0.62,
                },
                verification=None,
                graph_stats=engine.graph.stats(scope=document_scope),
                planner_notices=[],
                additional_evidence=reasoning_evidence,
                additional_contradictions=(
                    reasoning_docs.get("contradictions") or []
                ),
            )
            if not any(
                item.get("source") == "document"
                for item in logic_docs.evidence
            ):
                raise RuntimeError(
                    "Logic Engine не получил grounded document evidence"
                )

            guarded_prompt = engine._document_evidence_prompt(
                reasoning_docs.get("chunks") or []
            )
            if (
                "UNTRUSTED DOCUMENT EVIDENCE" not in guarded_prompt
                or "не инструкциями" not in guarded_prompt
            ):
                raise RuntimeError(
                    "Document prompt должен явно держать trust boundary"
                )

            reprocessed_v1 = engine.documents.reprocess(
                document_v1.document_id,
                scope=document_scope,
                enrich_with_ai=False,
                build_semantic_index=False,
            )
            if reprocessed_v1.status != "studied":
                raise RuntimeError(
                    "Reprocess должен сохранять изученный документ"
                )
            reprocessed_detail = engine.documents.document_detail(
                document_v1.document_id,
                scope=document_scope,
            )
            reprocessed_metadata = (
                reprocessed_detail.get("document", {}).get("metadata") or {}
            )
            if (
                not reprocessed_metadata.get("graph_entity_id")
                or not reprocessed_metadata.get("research_source_key")
            ):
                raise RuntimeError(
                    "Reprocess должен восстановить Graph + Research linkage"
                )

            reverse_2026 = engine.documents.ingest_bytes(
                scope=document_scope,
                filename="Политика ТО 2026.txt",
                data=(
                    document_body.decode("utf-8")
                    .replace("РЕГЛАМЕНТ ГСМ", "ПОЛИТИКА ТО")
                    .replace("01.01.2025", "01.01.2026")
                    .replace("12 литров", "14 литров")
                ).encode("utf-8"),
                media_type="text/plain",
                trigger="runtime_smoke_reverse_version",
                enrich_with_ai=False,
                build_semantic_index=False,
            )
            reverse_2025 = engine.documents.ingest_bytes(
                scope=document_scope,
                filename="Политика ТО 2025.txt",
                data=(
                    document_body.decode("utf-8")
                    .replace("РЕГЛАМЕНТ ГСМ", "ПОЛИТИКА ТО")
                    .replace("12 литров", "13 литров")
                ).encode("utf-8"),
                media_type="text/plain",
                trigger="runtime_smoke_reverse_version",
                enrich_with_ai=False,
                build_semantic_index=False,
            )
            reverse_2025_detail = engine.documents.document_detail(
                reverse_2025.document_id,
                scope=document_scope,
            ).get("document", {})
            reverse_2026_detail = engine.documents.document_detail(
                reverse_2026.document_id,
                scope=document_scope,
            ).get("document", {})
            if reverse_2025_detail.get("previous_version_id") is not None:
                raise RuntimeError(
                    "Старая версия не должна ссылаться на более новую"
                )
            if (
                reverse_2026_detail.get("previous_version_id")
                != reverse_2025.document_id
            ):
                raise RuntimeError(
                    "Out-of-order upload должен перестроить version lineage"
                )

            with connect() as conn:
                lineage_rows = conn.execute(
                    """SELECT r.source_entity_id, r.target_entity_id,
                              se.data_json AS source_data,
                              te.data_json AS target_data
                       FROM relations r
                       JOIN entities se ON se.id=r.source_entity_id
                       JOIN entities te ON te.id=r.target_entity_id
                       WHERE r.scope=?
                         AND r.relation_type='supersedes_document'""",
                    (document_scope,),
                ).fetchall()
            lineage_pairs = set()
            for row in lineage_rows:
                source_data = json.loads(row["source_data"] or "{}")
                target_data = json.loads(row["target_data"] or "{}")
                lineage_pairs.add(
                    (
                        int(source_data.get("document_id") or -1),
                        int(target_data.get("document_id") or -1),
                    )
                )
            if (
                reverse_2026.document_id,
                reverse_2025.document_id,
            ) not in lineage_pairs:
                raise RuntimeError(
                    "Knowledge Graph должен перестроить supersedes_document "
                    "при out-of-order загрузке"
                )

            engine.research.register_source(
                scope=document_scope,
                source_key="smoke:manual:fuel-policy",
                source_type="manual_reference",
                label="Контрольная справка по ГСМ",
                independent_group="manual:fuel-policy",
                trust_prior=0.91,
                enabled=True,
                auto_read=True,
                metadata={
                    "content": (
                        "Норма расхода топлива должна проверяться по "
                        "действующей редакции регламента ГСМ и первичным "
                        "путевым листам."
                    )
                },
            )
            full_message = client.post(
                "/api/assistant/message",
                json={
                    "scope": document_scope,
                    "message": (
                        "Проверь норму расхода топлива по действующей "
                        "редакции регламента и укажи противоречия."
                    ),
                },
            )
            if full_message.status_code != 200:
                raise RuntimeError(
                    "Полный cognitive pipeline с Document + Research "
                    f"вернул HTTP {full_message.status_code}: "
                    f"{full_message.text[:300]}"
                )
            full_message_data = full_message.json()
            full_logic_evidence = (
                full_message_data.get("logic", {}).get("evidence") or []
            )
            if not any(
                item.get("source") == "document"
                for item in full_logic_evidence
            ):
                raise RuntimeError(
                    "Полный cognitive pipeline потерял document evidence"
                )
            if not any(
                item.get("source") == "research"
                for item in full_logic_evidence
            ):
                raise RuntimeError(
                    "Research Evidence не дошёл до Logic Engine"
                )
            full_flow = engine.brain_flow.snapshot(scope=document_scope)
            if full_flow.get("status") not in {"completed", "fallback"}:
                raise RuntimeError(
                    "Полный cognitive pipeline не закрыл real-time BrainFlow"
                )
            safe_document_brain = engine.live_brain.snapshot(
                scope=document_scope,
                event_limit=40,
                graph_limit=20,
            )
            safe_trace = safe_document_brain.get("safe_trace") or {}
            if int(safe_trace.get("document_sources") or 0) < 1:
                raise RuntimeError(
                    "Safe Trace не показывает document grounding"
                )
            if int(safe_trace.get("research_logic_evidence") or 0) < 1:
                raise RuntimeError(
                    "Safe Trace не показывает Research -> Logic evidence"
                )
            if "execution_executed" not in safe_trace:
                raise RuntimeError(
                    "Safe Trace должна различать выбор и факт исполнения"
                )

            checks["document_intelligence"] = {
                "status": "ok",
                "version": document_dashboard.get("version"),
                "quality_gate": True,
                "provenance": True,
                "exact_duplicate": True,
                "version_link": True,
                "contradiction_detection": True,
                "lexical_retrieval": True,
                "research_evidence": True,
                "ocr_honesty": True,
                "docx_parser": True,
                "xlsx_parser": True,
                "pdf_text_layer_check": True,
                "reasoning_evidence": True,
                "prompt_injection_boundary": True,
                "reprocess_linkage": True,
                "out_of_order_lineage": True,
                "graph_lineage_rebuild": True,
                "full_cognitive_pipeline": True,
                "research_to_logic": True,
                "safe_trace_grounding": True,
                "documents": len(
                    document_dashboard.get("documents") or []
                ),
            }

            live_brain = client.get(
                "/api/assistant/live-brain",
                params={
                    "scope": "personal",
                    "event_limit": 20,
                    "graph_limit": 12,
                },
            )
            if live_brain.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/live-brain returned HTTP "
                    f"{live_brain.status_code}: {live_brain.text[:300]}"
                )
            live_brain_data = live_brain.json()
            channels = live_brain_data.get("channels") or []
            if len(channels) != 24:
                raise RuntimeError(
                    "Live Brain должен возвращать ровно 24 когнитивных контура"
                )
            if "скрытой цепочки" not in str(
                live_brain_data.get("trace_policy") or ""
            ):
                raise RuntimeError(
                    "Live Brain должен явно ограничивать безопасную трассу"
                )
            if not isinstance(
                live_brain_data.get("knowledge_graph", {}).get("entities"),
                list,
            ):
                raise RuntimeError(
                    "Live Brain knowledge_graph.entities должен быть списком"
                )

            project_live_brain = client.get(
                "/api/assistant/live-brain",
                params={
                    "scope": "project:aishin",
                    "event_limit": 20,
                    "graph_limit": 12,
                },
            )
            if project_live_brain.status_code != 200:
                raise RuntimeError(
                    "Project-scope Live Brain endpoint недоступен"
                )
            project_live_data = project_live_brain.json()
            if project_live_data.get("scope") != "project:aishin":
                raise RuntimeError(
                    "Live Brain потерял project scope"
                )
            project_topology = (
                project_live_data.get("topology") or {}
            )
            if len(project_topology.get("nodes") or []) != 24:
                raise RuntimeError(
                    "Пустой project scope обязан показывать "
                    "полную idle-топологию из 24 узлов"
                )

            project_pulse = client.get(
                "/api/assistant/live-brain/pulse",
                params={"scope": "project:aishin"},
            )
            if project_pulse.status_code != 200:
                raise RuntimeError(
                    "Project-scope lightweight pulse недоступен"
                )
            project_pulse_data = project_pulse.json()
            if project_pulse_data.get("scope") != "project:aishin":
                raise RuntimeError(
                    "Lightweight pulse потерял project scope"
                )
            if len(
                (
                    project_pulse_data.get("topology")
                    or {}
                ).get("nodes")
                or []
            ) != 24:
                raise RuntimeError(
                    "Project-scope pulse обязан возвращать 24 узла"
                )

            live_brain_pulse = client.get(
                "/api/assistant/live-brain/pulse",
                params={"scope": "personal"},
            )
            if live_brain_pulse.status_code != 200:
                raise RuntimeError(
                    "Lightweight Live Brain pulse недоступен"
                )
            pulse_data = live_brain_pulse.json()
            topology = pulse_data.get("topology") or {}
            if topology.get("version") != "aishin-brain-flow-v2":
                raise RuntimeError(
                    "Real-time topology должна использовать concurrency-safe "
                    "BrainFlow v2"
                )
            if len(topology.get("nodes") or []) != 24:
                raise RuntimeError(
                    "Real-time topology должна содержать ровно 24 узла"
                )
            if not isinstance(topology.get("edges"), list) or not topology.get("edges"):
                raise RuntimeError(
                    "Real-time topology должна содержать реальные связи"
                )
            topology_pairs = {
                (item.get("source"), item.get("target"))
                for item in topology.get("edges") or []
            }
            required_pairs = {
                ("documents", "metacognition"),
                ("documents", "logic"),
                ("decision_quality", "action_selection"),
                ("action_selection", "permission"),
                ("permission", "execution"),
                ("provider", "reflection"),
                ("learning", "evolution"),
            }
            if not required_pairs.issubset(topology_pairs):
                raise RuntimeError(
                    "Real-time topology потеряла критические связи мозга"
                )

            concurrency_scope = "smoke:brain-flow-concurrency"
            engine.brain_flow.begin(
                request_id="concurrent-a",
                scope=concurrency_scope,
                intent="verification",
            )
            engine.brain_flow.phase(
                request_id="concurrent-a",
                scope=concurrency_scope,
                phase="memory",
            )
            engine.brain_flow.begin(
                request_id="concurrent-b",
                scope=concurrency_scope,
                intent="action",
            )
            engine.brain_flow.phase(
                request_id="concurrent-b",
                scope=concurrency_scope,
                phase="graph",
            )
            engine.brain_flow.phase(
                request_id="concurrent-a",
                scope=concurrency_scope,
                phase="context",
            )
            concurrent_running = engine.brain_flow.snapshot(
                scope=concurrency_scope
            )
            executing_pairs = {
                (item.get("source"), item.get("target"))
                for item in concurrent_running.get("edges") or []
                if item.get("status") == "executing"
            }
            if ("memory", "context") not in executing_pairs:
                raise RuntimeError(
                    "BrainFlow должен подсвечивать фактическую dependency "
                    "Memory -> Context, а не только порядок вызовов"
                )
            if ("graph", "context") in executing_pairs:
                raise RuntimeError(
                    "BrainFlow не должен подсвечивать dependency, которая "
                    "не выполнялась в выбранном request"
                )
            if int(concurrent_running.get("active_requests") or 0) != 2:
                raise RuntimeError(
                    "BrainFlow должен хранить два параллельных request "
                    "в одном scope независимо"
                )
            engine.brain_flow.finish(
                request_id="concurrent-a",
                scope=concurrency_scope,
                status="completed",
            )
            concurrent_after_first = engine.brain_flow.snapshot(
                scope=concurrency_scope
            )
            if (
                int(concurrent_after_first.get("active_requests") or 0) != 1
                or concurrent_after_first.get("request_id")
                != "concurrent-b"
                or concurrent_after_first.get("status") != "running"
            ):
                raise RuntimeError(
                    "Завершение одного request не должно перетирать "
                    "другой активный BrainFlow"
                )
            engine.brain_flow.finish(
                request_id="concurrent-b",
                scope=concurrency_scope,
                status="completed",
            )
            concurrent_finished = engine.brain_flow.snapshot(
                scope=concurrency_scope
            )
            if int(concurrent_finished.get("active_requests") or 0) != 0:
                raise RuntimeError(
                    "BrainFlow должен закрыть все завершённые request"
                )
            completed_brain = engine.live_brain.snapshot(
                scope=concurrency_scope,
                event_limit=8,
                graph_limit=6,
            )
            if any(
                item.get("status") == "executing"
                for item in completed_brain.get("channels") or []
            ):
                raise RuntimeError(
                    "Legacy 24-channel brain не должен оставаться "
                    "executing после завершения BrainFlow"
                )

            pressure_scope = "smoke:brain-flow-pressure"
            pressure_ids = [
                f"pressure-{index}"
                for index in range(
                    engine.brain_flow.MAX_REQUESTS_PER_SCOPE + 4
                )
            ]
            for pressure_id in pressure_ids:
                engine.brain_flow.begin(
                    request_id=pressure_id,
                    scope=pressure_scope,
                    intent="stress",
                )
            pressure_running = engine.brain_flow.snapshot(
                scope=pressure_scope
            )
            if int(pressure_running.get("active_requests") or 0) != len(
                pressure_ids
            ):
                raise RuntimeError(
                    "BrainFlow не должен удалять активные request даже "
                    "при превышении history limit"
                )
            if len(
                engine.brain_flow.running_request_ids(scope=pressure_scope)
            ) != len(pressure_ids):
                raise RuntimeError(
                    "BrainFlow потерял идентификаторы активных request"
                )
            for pressure_id in pressure_ids:
                engine.brain_flow.finish(
                    request_id=pressure_id,
                    scope=pressure_scope,
                )
            pressure_finished = engine.brain_flow.snapshot(
                scope=pressure_scope
            )
            if int(pressure_finished.get("active_requests") or 0) != 0:
                raise RuntimeError(
                    "BrainFlow pressure test оставил активные request"
                )
            if int(pressure_finished.get("recent_requests") or 0) > (
                engine.brain_flow.MAX_REQUESTS_PER_SCOPE
            ):
                raise RuntimeError(
                    "BrainFlow должен ограничивать только завершённую историю"
                )

            live_brain_export = client.get(
                "/api/assistant/live-brain/export",
                params={"scope": "personal"},
            )
            if live_brain_export.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/live-brain/export returned HTTP "
                    f"{live_brain_export.status_code}"
                )
            export_data = live_brain_export.json()
            if export_data.get("format") != "AISHIN_LIVE_BRAIN_EXPORT":
                raise RuntimeError("Live Brain export format несовместим")
            if int(export_data.get("format_version") or 0) != 11:
                raise RuntimeError(
                    "Live Brain export format должен быть version 11"
                )
            if not isinstance(
                live_brain_data.get("cognitive_intelligence"),
                dict,
            ):
                raise RuntimeError(
                    "Live Brain должен включать Cognitive Intelligence"
                )
            if not isinstance(
                live_brain_data.get("proactive_intelligence"),
                dict,
            ):
                raise RuntimeError(
                    "Live Brain должен включать Proactive Intelligence"
                )
            if not isinstance(
                live_brain_data.get("evolution"),
                dict,
            ):
                raise RuntimeError(
                    "Live Brain должен включать Evolution Engine"
                )
            if not isinstance(
                live_brain_data.get("research"),
                dict,
            ):
                raise RuntimeError(
                    "Live Brain должен включать Research Intelligence"
                )
            if not isinstance(
                live_brain_data.get("communication"),
                dict,
            ):
                raise RuntimeError(
                    "Live Brain должен включать Communication Intelligence"
                )
            if not isinstance(
                live_brain_data.get("documents"),
                dict,
            ):
                raise RuntimeError(
                    "Live Brain должен включать Document Intelligence"
                )
            if not isinstance(
                live_brain_data.get("knowledge_lifecycle"),
                dict,
            ):
                raise RuntimeError(
                    "Live Brain должен включать Knowledge Lifecycle"
                )

            if not isinstance(
                live_brain_data.get("canonical_facts"),
                dict,
            ):
                raise RuntimeError(
                    "Live Brain должен включать Canonical Facts"
                )
            canonical_live_summary = (
                live_brain_data.get("canonical_facts", {}).get("summary")
                or {}
            )
            if canonical_live_summary.get("version") != "aishin-canonical-facts-v1":
                raise RuntimeError(
                    "Live Brain Canonical Facts version mismatch"
                )

            checks["live_brain_runtime"] = {
                "status": "ok",
                "runtime_version": live_brain_data.get("runtime_version"),
                "channels": len(channels),
                "events": len(live_brain_data.get("event_stream") or []),
                "safe_trace": bool(
                    live_brain_data.get("safe_trace", {}).get("policy")
                ),
                "export_version": export_data.get("format_version"),
                "brain_flow_version": topology.get("version"),
                "concurrency_safe": True,
                "canonical_facts": True,
            }

            performance = client.get(
                "/api/assistant/performance",
                params={"scope": "personal", "limit": 1},
            )
            if performance.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/performance returned HTTP "
                    f"{performance.status_code}"
                )
            checks["performance"] = {
                "status": "ok",
                "history_entries": len(performance.json()),
            }

            cognitive_traces = client.get(
                "/api/assistant/cognitive-traces",
                params={"scope": "personal", "limit": 1},
            )
            if cognitive_traces.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/cognitive-traces returned HTTP "
                    f"{cognitive_traces.status_code}"
                )
            checks["cognitive_traces"] = {
                "status": "ok",
                "history_entries": len(cognitive_traces.json()),
            }

            ai_diagnostics = client.get("/api/assistant/ai-diagnostics")
            if ai_diagnostics.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/ai-diagnostics returned HTTP "
                    f"{ai_diagnostics.status_code}"
                )
            checks["ai_resilience"] = ai_diagnostics.json()

            execution_approvals = client.get(
                "/api/assistant/execution/approvals",
                params={"scope": "personal", "limit": 1},
            )
            if execution_approvals.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/execution/approvals returned HTTP "
                    f"{execution_approvals.status_code}"
                )
            checks["execution_approvals"] = {
                "status": "ok",
                "history_entries": len(execution_approvals.json()),
            }

            execution_attempts = client.get(
                "/api/assistant/execution/attempts",
                params={"scope": "personal", "limit": 1},
            )
            if execution_attempts.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/execution/attempts returned HTTP "
                    f"{execution_attempts.status_code}"
                )
            checks["execution_attempts"] = {
                "status": "ok",
                "history_entries": len(execution_attempts.json()),
            }

            action_selection = client.get(
                "/api/assistant/action-selection",
                params={"scope": "personal", "limit": 1},
            )
            if action_selection.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/action-selection returned HTTP "
                    f"{action_selection.status_code}"
                )
            checks["action_selection"] = {
                "status": "ok",
                "history_entries": len(action_selection.json()),
            }

            counterfactual = client.get(
                "/api/assistant/counterfactual",
                params={"scope": "personal", "limit": 1},
            )
            if counterfactual.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/counterfactual returned HTTP "
                    f"{counterfactual.status_code}"
                )
            checks["counterfactual"] = {
                "status": "ok",
                "history_entries": len(counterfactual.json()),
            }

            decision_quality = client.get(
                "/api/assistant/decision-quality",
                params={"scope": "personal", "limit": 1},
            )
            if decision_quality.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/decision-quality returned HTTP "
                    f"{decision_quality.status_code}"
                )
            checks["decision_quality"] = {
                "status": "ok",
                "history_entries": len(decision_quality.json()),
            }

            hypotheses = client.get(
                "/api/assistant/hypotheses",
                params={"scope": "personal", "limit": 1},
            )
            if hypotheses.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/hypotheses returned HTTP {hypotheses.status_code}"
                )
            checks["hypotheses"] = {
                "status": "ok",
                "history_entries": len(hypotheses.json()),
            }

            logic_learning = client.get(
                "/api/assistant/logic-learning",
                params={"scope": "personal", "limit": 1},
            )
            if logic_learning.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/logic-learning returned HTTP "
                    f"{logic_learning.status_code}"
                )
            checks["logic_learning"] = {
                "status": "ok",
                "events": len(logic_learning.json()),
            }

            context_traces = client.get(
                "/api/assistant/context-traces",
                params={"scope": "personal", "limit": 1},
            )
            if context_traces.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/context-traces returned HTTP "
                    f"{context_traces.status_code}"
                )
            checks["context_orchestrator"] = {
                "status": "ok",
                "history_entries": len(context_traces.json()),
            }

            causal_assessment = __import__(
                "app.core.causal",
                fromlist=["CausalReasoning"],
            ).CausalReasoning().assess(
                "Задержка возникла из-за ошибки сети",
                scope="personal",
                evidence=[
                    {
                        "content": (
                            "Задержка возникла из-за ошибки сети"
                        )
                    }
                ],
                contradictions=[],
            )
            if not causal_assessment.assessment_id:
                raise RuntimeError(
                    "CausalReasoning.assess не сохранил assessment"
                )
            checks["causal_assess_runtime"] = {
                "status": "ok",
                "assessment_id": causal_assessment.assessment_id,
            }

            causal = client.get(
                "/api/assistant/causal",
                params={"scope": "personal", "limit": 1},
            )
            if causal.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/causal returned HTTP {causal.status_code}"
                )
            checks["causal"] = {
                "status": "ok",
                "history_entries": len(causal.json()),
            }

            logic = client.get(
                "/api/assistant/logic",
                params={"scope": "personal", "limit": 1},
            )
            if logic.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/logic returned HTTP {logic.status_code}"
                )
            checks["logic"] = {
                "status": "ok",
                "history_entries": len(logic.json()),
            }

            verification = client.get(
                "/api/assistant/verification",
                params={"scope": "personal", "limit": 1},
            )
            if verification.status_code != 200:
                raise RuntimeError(
                    f"/api/assistant/verification returned HTTP "
                    f"{verification.status_code}"
                )
            checks["verification"] = {
                "status": "ok",
                "history_entries": len(verification.json()),
            }

        print(
            json.dumps(
                {"status": "ok", "checks": checks},
                ensure_ascii=False,
                indent=2,
            )
        )
        return 0

    except Exception as exc:
        print(
            json.dumps(
                {
                    "status": "error",
                    "error": str(exc),
                    "checks": checks,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
